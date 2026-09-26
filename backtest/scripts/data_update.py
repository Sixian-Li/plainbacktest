#!/usr/bin/env python3
"""Operate the Quant market-data layer without silently mutating approved data.

Commands:
  status            Summarize freshness, source evidence, credentials, and imports.
  check             Run quick or deep integrity checks.
  shadow-update     Fetch isolated SPY-current-member update candidates.
  import-purchased  Inventory a new purchase; copy only with explicit --apply.
  audit-nasdaq100   Audit purchased Nasdaq-100 evidence; optionally build a candidate.
  audit-sp500       Reaudit purchased S&P 500 evidence into isolated pending_review products.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

import exchange_calendars as xcals
from openpyxl import load_workbook

WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
BACKTEST_ROOT = WORKSPACE_ROOT / "backtest"
if str(BACKTEST_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKTEST_ROOT))

from scripts.validate_sp500_market_sources import (
    FIELDS,
    TIINGO_URL,
    TWELVE_URL,
    ProviderError,
    archive_payload,
    atomic_write,
    compare_series,
    csv_bytes,
    curl_json,
    json_bytes,
    parse_tiingo,
    parse_twelve,
    percentile,
    sha256_bytes,
    tiingo_symbol,
    validate_rows,
)
from quantkit.nasdaq100_data import PointInTimeDataError, audit_and_build


CONFIG_PATH = WORKSPACE_ROOT / "data/data_update_registry.json"
NASDAQ100_REGISTRY_PATH = WORKSPACE_ROOT / "data/nasdaq100_history_registry.json"
SP500_CANDIDATE_REGISTRY_PATH = WORKSPACE_ROOT / "data/sp500_candidate_registry.json"
BASELINE_REGISTRY_PATH = WORKSPACE_ROOT / "data/sp500_history_registry.json"
BASELINE_MANIFEST_PATH = WORKSPACE_ROOT / "data/processed/universes/sp500/manifest.json"
BASELINE_QUALITY_PATH = WORKSPACE_ROOT / "data/processed/universes/sp500/quality_report.json"
SOURCE_VALIDATION_LATEST = WORKSPACE_ROOT / "data/processed/validations/sp500_market_sources/latest.json"
CANONICAL_COLUMNS = ("date", "symbol", "open", "high", "low", "close", "volume")
SUPPORTED_INSPECTION_SUFFIXES = {".csv", ".json", ".zip", ".xlsx", ".xls"}
KEYCHAIN_SERVICES = {
    "TWELVE_DATA_API_KEY": "quant.twelvedata.api_key",
    "TIINGO_API_TOKEN": "quant.tiingo.api_token",
}


class UpdateError(RuntimeError):
    """A data-update operation cannot safely continue."""


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def credential(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if value:
        return value
    service = KEYCHAIN_SERVICES.get(name)
    account = os.environ.get("USER", "")
    if sys.platform == "darwin" and service and account:
        result = subprocess.run(
            ["security", "find-generic-password", "-a", account, "-s", service, "-w"],
            capture_output=True,
            text=True,
        )
        if result.returncode == 0:
            return result.stdout.strip()
    return ""


def config() -> dict[str, Any]:
    payload = load_json(CONFIG_PATH)
    if payload.get("production_writes_enabled") is not False:
        raise UpdateError("This implementation requires production_writes_enabled=false")
    return payload


def security_identity_reviews(cfg: dict[str, Any]) -> dict[str, dict[str, Any]]:
    path_text = cfg.get("security_identity_reviews")
    if not path_text:
        return {}
    payload = load_json(WORKSPACE_ROOT / path_text)
    reviews = {item["symbol"]: item for item in payload.get("reviews", [])}
    current_path = WORKSPACE_ROOT / cfg["universe"]["current_constituents"]
    with current_path.open(newline="", encoding="utf-8") as file:
        current = {item["symbol"]: item for item in csv.DictReader(file)}
    for symbol, review in reviews.items():
        evidence_path = WORKSPACE_ROOT / review["official_snapshot_path"]
        if evidence_path.resolve() != current_path.resolve():
            raise UpdateError(f"Identity review for {symbol} references an unexpected official snapshot path")
        if sha256_file(evidence_path) != review["official_snapshot_sha256"]:
            raise UpdateError(f"Identity review for {symbol} is stale after the official snapshot changed")
        row = current.get(symbol)
        if not row:
            raise UpdateError(f"Identity review symbol is absent from the current snapshot: {symbol}")
        if row.get("identifier") != review.get("official_identifier"):
            raise UpdateError(f"Identity review identifier mismatch for {symbol}")
        if row.get("company_name") != review.get("official_company_name"):
            raise UpdateError(f"Identity review company-name mismatch for {symbol}")
    return reviews


def corporate_action_reviews(cfg: dict[str, Any]) -> dict[str, dict[str, Any]]:
    path_text = cfg.get("corporate_action_reviews")
    if not path_text:
        return {}
    payload = load_json(WORKSPACE_ROOT / path_text)
    reviews = {item["symbol"]: item for item in payload.get("reviews", [])}
    for symbol, review in reviews.items():
        if review.get("decision") != "documented_public_corporate_action":
            raise UpdateError(f"Unsupported corporate-action decision for {symbol}")
        dates = review.get("covered_return_outlier_dates", [])
        if not dates or len(dates) != len(set(dates)):
            raise UpdateError(f"Corporate-action review dates are empty or duplicated for {symbol}")
        for value in dates:
            date.fromisoformat(value)
        if not review.get("events") or any(not event.get("url") for event in review["events"]):
            raise UpdateError(f"Corporate-action review has no source evidence for {symbol}")
    return reviews


def read_current_constituents(cfg: dict[str, Any]) -> list[dict[str, str]]:
    path = WORKSPACE_ROOT / cfg["universe"]["current_constituents"]
    with path.open(newline="", encoding="utf-8") as file:
        rows = list(csv.DictReader(file))
    symbols = [row["symbol"] for row in rows]
    if len(rows) != cfg["universe"]["expected_equities"]:
        raise UpdateError(f"Expected {cfg['universe']['expected_equities']} current equities; found {len(rows)}")
    if len(symbols) != len(set(symbols)):
        raise UpdateError("Current constituent snapshot contains duplicate symbols")
    if len({row["as_of_date"] for row in rows}) != 1:
        raise UpdateError("Current constituent snapshot contains multiple as-of dates")
    return rows


def read_last_csv_record(path: Path) -> dict[str, str]:
    with path.open(newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        if tuple(reader.fieldnames or ()) != CANONICAL_COLUMNS:
            raise UpdateError(f"Unexpected canonical schema: {path}")
        last: dict[str, str] | None = None
        for row in reader:
            last = row
    if last is None:
        raise UpdateError(f"Empty canonical price file: {path}")
    return last


def read_canonical_from(path: Path, start: date) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        if tuple(reader.fieldnames or ()) != CANONICAL_COLUMNS:
            raise UpdateError(f"Unexpected canonical schema: {path}")
        for raw in reader:
            if date.fromisoformat(raw["date"]) < start:
                continue
            row = {"date": raw["date"]}
            for field in FIELDS:
                value = float(raw[field])
                if not math.isfinite(value):
                    raise UpdateError(f"Non-finite {field} in {path} on {raw['date']}")
                row[field] = value
            rows.append(row)
    validate_rows(rows, "approved_canonical", path.stem)
    return rows


def completed_session_on_or_before(
    end: date | None = None,
    *,
    now_utc: datetime | None = None,
) -> date:
    calendar = xcals.get_calendar("XNYS")
    now = now_utc or datetime.now(timezone.utc)
    if now.tzinfo is None:
        raise UpdateError("now_utc must be timezone-aware")
    requested_end = end or now.date()
    search_end = min(requested_end, now.date())
    sessions = calendar.sessions_in_range(search_end - timedelta(days=14), search_end)
    completed = [
        session
        for session in sessions
        if session.date() <= requested_end and calendar.session_close(session).to_pydatetime() <= now
    ]
    if not completed:
        raise UpdateError(f"Could not determine a completed XNYS session on or before {requested_end}")
    return completed[-1].date()


def expected_sessions(start: date, end: date) -> list[str]:
    if start > end:
        return []
    calendar = xcals.get_calendar("XNYS")
    return [session.date().isoformat() for session in calendar.sessions_in_range(start, end)]


def verify_pointer(pointer_path: Path) -> tuple[dict[str, Any] | None, list[str]]:
    if not pointer_path.is_file():
        return None, [f"missing pointer: {pointer_path.relative_to(WORKSPACE_ROOT)}"]
    pointer = load_json(pointer_path)
    report_path_text = pointer.get("report_json")
    if not report_path_text:
        return pointer, [f"pointer has no report_json: {pointer_path.relative_to(WORKSPACE_ROOT)}"]
    report_path = WORKSPACE_ROOT / report_path_text
    if not report_path.is_file():
        return pointer, [f"pointer target is missing: {report_path_text}"]
    expected = pointer.get("report_sha256")
    if expected and sha256_file(report_path) != expected:
        return pointer, [f"pointer target hash mismatch: {report_path_text}"]
    return pointer, []


def verify_shadow_report_artifacts(report: dict[str, Any]) -> tuple[int, int]:
    candidate_files = report.get("candidate_files", {})
    if len(candidate_files) != int(report.get("summary", {}).get("symbols", -1)):
        raise UpdateError(
            f"Shadow candidate-file count mismatch: files={len(candidate_files)}; "
            f"symbols={report.get('summary', {}).get('symbols')}"
        )
    verified_candidates = 0
    for symbol, item in candidate_files.items():
        path = WORKSPACE_ROOT / item["path"]
        if not path.is_file() or sha256_file(path) != item["sha256"]:
            raise UpdateError(f"Shadow candidate SHA256 mismatch for {symbol}: {item['path']}")
        with path.open(newline="", encoding="utf-8") as file:
            rows = sum(1 for _row in csv.DictReader(file))
        if rows != int(item["rows"]):
            raise UpdateError(f"Shadow candidate row-count mismatch for {symbol}: {rows} != {item['rows']}")
        verified_candidates += 1
    source_archives: dict[str, dict[str, Any]] = {}
    for batch in report.get("twelve_data", {}).get("batches", []):
        archive = batch.get("archive") or batch.get("source_archive")
        if archive:
            source_archives[archive["path"]] = archive
    for item in report.get("tiingo", {}).get("metadata", {}).values():
        archive = item.get("archive")
        if archive:
            source_archives[archive["path"]] = archive
    for item in source_archives.values():
        path = WORKSPACE_ROOT / item["path"]
        if not path.is_file() or sha256_file(path) != item["sha256"]:
            raise UpdateError(f"Shadow source archive SHA256 mismatch: {item['path']}")
    return verified_candidates, len(source_archives)


def qualify_shadow_report(report: dict[str, Any], cfg: dict[str, Any]) -> tuple[bool, list[str]]:
    campaign = cfg["shadow_validation_campaign"]
    expected = int(cfg["universe"]["expected_equities"])
    required_cross_checks = int(campaign["required_tiingo_cross_checks_per_session"])
    allowed_statuses = {"ready_candidate"}
    if campaign.get("allow_verified_shadow_bootstrap_candidate"):
        allowed_statuses.add("ready_bootstrap_candidate")
    reasons: list[str] = []
    scope = report.get("scope", {})
    summary = report.get("summary", {})
    symbols = report.get("symbols", [])
    if scope.get("kind") != "full_current_universe" or scope.get("selected_count") != expected:
        reasons.append("not_full_current_universe")
    if summary.get("symbols") != expected or len(symbols) != expected:
        reasons.append("symbol_count_mismatch")
    nonready = [item.get("symbol") for item in symbols if item.get("status") not in allowed_statuses]
    if nonready:
        reasons.append(f"nonready_symbols={','.join(str(item) for item in nonready[:10])}")
    if int(summary.get("fetch_failed", 0)) != 0:
        reasons.append("fetch_failures")
    twelve_failures = report.get("twelve_data", {}).get("symbol_failures", [])
    if twelve_failures:
        reasons.append("twelve_symbol_failures")
    requested = report.get("tiingo", {}).get("requested_symbols", [])
    requested_set = set(requested)
    tiingo_failures = report.get("tiingo", {}).get("failures", [])
    if len(requested) < required_cross_checks:
        reasons.append("insufficient_tiingo_cross_checks")
    if tiingo_failures or int(summary.get("cross_check_passed", 0)) != len(requested):
        reasons.append("tiingo_cross_check_not_all_passed")
    warning_symbols = {
        item.get("symbol")
        for item in symbols
        if "historical_adjustment_basis_change" in item.get("warnings", [])
    }
    uncovered_warnings = sorted(symbol for symbol in warning_symbols if symbol not in requested_set)
    if uncovered_warnings:
        reasons.append(f"adjustment_warnings_not_cross_checked={','.join(uncovered_warnings[:10])}")
    bootstrap_symbols = {
        item.get("symbol") for item in symbols if item.get("status") == "ready_bootstrap_candidate"
    }
    unchecked_bootstraps = sorted(symbol for symbol in bootstrap_symbols if symbol not in requested_set)
    if unchecked_bootstraps:
        reasons.append(f"bootstrap_candidates_not_cross_checked={','.join(unchecked_bootstraps[:10])}")
    if int(summary.get("production_rows_written", -1)) != 0 or report.get("production_writes_enabled") is not False:
        reasons.append("shadow_write_boundary_failed")
    return not reasons, reasons


def shadow_campaign_status(cfg: dict[str, Any]) -> dict[str, Any]:
    root = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"]
    runs: list[dict[str, Any]] = []
    for report_path in sorted(root.glob("run_*/report.json")) if root.is_dir() else []:
        try:
            report = load_json(report_path)
            if report.get("scope", {}).get("kind") != "full_current_universe":
                continue
            qualified, reasons = qualify_shadow_report(report, cfg)
            runs.append(
                {
                    "run_id": report.get("run_id"),
                    "desired_end_session": report.get("desired_end_session"),
                    "generated_at_utc": report.get("generated_at_utc"),
                    "qualified": qualified,
                    "reasons": reasons,
                    "report_json": str(report_path.relative_to(WORKSPACE_ROOT)),
                }
            )
        except (OSError, ValueError, json.JSONDecodeError):
            continue
    qualified_by_session: dict[str, dict[str, Any]] = {}
    for item in runs:
        session = item.get("desired_end_session")
        if not session or not item["qualified"]:
            continue
        prior = qualified_by_session.get(session)
        if prior is None or str(item.get("generated_at_utc")) > str(prior.get("generated_at_utc")):
            qualified_by_session[session] = item
    sessions = sorted(qualified_by_session)
    target = int(cfg["shadow_validation_campaign"]["target_consecutive_completed_sessions"])
    streak_sessions: list[str] = []
    if sessions:
        calendar = xcals.get_calendar("XNYS")
        start = date.fromisoformat(sessions[0])
        end = date.fromisoformat(sessions[-1])
        expected = [session.date().isoformat() for session in calendar.sessions_in_range(start, end)]
        expected_position = {session: index for index, session in enumerate(expected)}
        for session in sessions:
            if not streak_sessions:
                streak_sessions = [session]
                continue
            prior = streak_sessions[-1]
            if expected_position.get(session) == expected_position.get(prior, -2) + 1:
                streak_sessions.append(session)
            else:
                streak_sessions = [session]
    return {
        "target_consecutive_completed_sessions": target,
        "qualified_distinct_sessions": len(sessions),
        "current_consecutive_sessions": len(streak_sessions),
        "current_streak_sessions": streak_sessions,
        "remaining_sessions": max(0, target - len(streak_sessions)),
        "qualified_sessions": sessions,
        "qualifying_runs": [qualified_by_session[session] for session in sessions],
        "full_runs_observed": len(runs),
        "full_runs_requiring_review": [item for item in runs if not item["qualified"]],
        "complete": len(streak_sessions) >= target,
    }


def database_freshness(rows: list[dict[str, str]], cfg: dict[str, Any]) -> dict[str, Any]:
    price_root = WORKSPACE_ROOT / cfg["universe"]["approved_price_directory"]
    distribution: Counter[str] = Counter()
    missing: list[str] = []
    mismatched: list[str] = []
    for row in rows:
        instrument = row["vendor_instrument_id"]
        available = row["vendor_price_available"].lower() == "true"
        if not available or not instrument:
            missing.append(row["symbol"])
            continue
        path = price_root / f"{instrument}.csv"
        if not path.is_file():
            missing.append(row["symbol"])
            continue
        last = read_last_csv_record(path)
        distribution[last["date"]] += 1
        if last["symbol"] != instrument:
            mismatched.append(row["symbol"])
    return {
        "price_available": len(rows) - len(missing),
        "price_missing": missing,
        "last_date_distribution": dict(sorted(distribution.items())),
        "symbol_mismatches": mismatched,
    }


def tracked_asset_freshness(cfg: dict[str, Any]) -> dict[str, Any]:
    """Expose QQQ/SPY/VOO freshness without implying that they are update-enabled."""
    tracked = cfg.get("tracked_assets", {})
    registry_path = WORKSPACE_ROOT / tracked["source_registry"]
    registry = load_json(registry_path)
    datasets = {item["symbol"]: item for item in registry.get("datasets", [])}
    assets: list[dict[str, Any]] = []
    for symbol in tracked.get("symbols", []):
        dataset = datasets.get(symbol)
        if not dataset:
            assets.append({"symbol": symbol, "status": "unregistered", "latest_date": None})
            continue
        output = WORKSPACE_ROOT / dataset["output"]
        latest_date = read_last_csv_record(output)["date"] if output.is_file() else None
        assets.append(
            {
                "symbol": symbol,
                "registered_status": dataset.get("registered_status"),
                "role": dataset.get("role"),
                "path": dataset["output"],
                "latest_date": latest_date,
            }
        )
    return {
        "assets": assets,
        "required_for_current_research": tracked.get("required_for_current_research", []),
        "shadow_update_enabled": bool(tracked.get("shadow_update_enabled", False)),
    }


def build_status() -> dict[str, Any]:
    cfg = config()
    rows = read_current_constituents(cfg)
    freshness = database_freshness(rows, cfg)
    tracked_assets = tracked_asset_freshness(cfg)
    baseline_manifest = load_json(BASELINE_MANIFEST_PATH)
    source_validation, source_pointer_issues = verify_pointer(SOURCE_VALIDATION_LATEST)
    shadow_latest_path = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"] / "latest.json"
    shadow_latest, shadow_pointer_issues = verify_pointer(shadow_latest_path)
    import_registry_path = WORKSPACE_ROOT / cfg["purchased_import"]["registry"]
    imports = load_json(import_registry_path)
    nasdaq100_manifest_path = WORKSPACE_ROOT / cfg["nasdaq100_candidate"]["manifest"]
    nasdaq100_manifest = load_json(nasdaq100_manifest_path) if nasdaq100_manifest_path.is_file() else None
    sp500_candidate_manifest_path = WORKSPACE_ROOT / cfg["sp500_candidate"]["manifest"]
    sp500_candidate_manifest = (
        load_json(sp500_candidate_manifest_path) if sp500_candidate_manifest_path.is_file() else None
    )
    approved_dates = sorted(freshness["last_date_distribution"])
    earliest_approved = approved_dates[0] if approved_dates else None
    latest_approved = approved_dates[-1] if approved_dates else None
    latest_completed = completed_session_on_or_before().isoformat()
    warnings: list[str] = []
    if earliest_approved and earliest_approved < latest_completed:
        warnings.append(
            f"approved current-member price floor ends at {earliest_approved}; "
            f"latest completed XNYS session is {latest_completed}"
        )
    if len(approved_dates) > 1:
        warnings.append(
            f"approved current-member files have mixed end dates from {earliest_approved} to {latest_approved}"
        )
    if freshness["price_missing"]:
        warnings.append(f"missing approved prices for {', '.join(freshness['price_missing'])}")
    required_assets = set(tracked_assets["required_for_current_research"])
    for asset in tracked_assets["assets"]:
        if asset["symbol"] in required_assets and asset.get("latest_date") != latest_completed:
            warnings.append(
                f"tracked research asset {asset['symbol']} ends at {asset.get('latest_date')}; "
                f"latest completed XNYS session is {latest_completed}"
            )
    warnings.extend(source_pointer_issues)
    warnings.extend(shadow_pointer_issues)
    if source_validation:
        source_summary = source_validation.get("summary", {})
        sample = int(source_summary.get("sample_symbols", 0))
        all_three = int(source_summary.get("symbols_with_all_three_sources", 0))
        if sample and all_three < sample:
            warnings.append(f"latest source validation has three-source coverage for {all_three}/{sample} symbols")
    if shadow_latest:
        shadow_scope = shadow_latest.get("scope", {})
        if shadow_scope.get("kind") != "full_current_universe":
            warnings.append(
                "latest shadow update is not full-universe: "
                f"{shadow_scope.get('selected_count', 'unknown')}/{shadow_scope.get('current_universe_count', 'unknown')}"
            )
        shadow_end = shadow_latest.get("desired_end_session")
        if shadow_end and shadow_end < latest_completed:
            warnings.append(f"latest shadow update ends at {shadow_end}; latest completed session is {latest_completed}")
    return {
        "schema_version": 1,
        "generated_at_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "production_writes_enabled": cfg["production_writes_enabled"],
        "capabilities": cfg.get("capabilities", {}),
        "schedule": cfg["schedule"],
        "tracked_assets": tracked_assets,
        "universe": {
            "as_of_date": rows[0]["as_of_date"],
            "constituents": len(rows),
            **freshness,
            "earliest_approved_date": earliest_approved,
            "latest_approved_date": latest_approved,
            "latest_completed_session": latest_completed,
        },
        "baseline": {
            "build_id": baseline_manifest.get("build_id"),
            "status": baseline_manifest.get("baseline_status"),
            "instrument_count": baseline_manifest.get("summary", {}).get("instrument_count"),
        },
        "credentials": {name: bool(credential(name)) for name in KEYCHAIN_SERVICES},
        "latest_source_validation": source_validation,
        "latest_shadow_update": shadow_latest,
        "shadow_validation_campaign": shadow_campaign_status(cfg),
        "purchased_imports": {
            "batches": len(imports.get("batches", [])),
            "status_counts": dict(Counter(item.get("status", "unknown") for item in imports.get("batches", []))),
        },
        "nasdaq100_candidate": {
            "present": nasdaq100_manifest is not None,
            "status": (nasdaq100_manifest or {}).get("status", "not_built"),
            "approved": (nasdaq100_manifest or {}).get("approved", False),
            "build_id": (nasdaq100_manifest or {}).get("build_id"),
            "coverage": (nasdaq100_manifest or {}).get("coverage"),
        },
        "sp500_candidate": {
            "present": sp500_candidate_manifest is not None,
            "status": (sp500_candidate_manifest or {}).get("status", "not_built"),
            "approved": (sp500_candidate_manifest or {}).get("approved", False),
            "build_id": (sp500_candidate_manifest or {}).get("build_id"),
            "coverage": (sp500_candidate_manifest or {}).get("coverage"),
            "existing_approved_baseline_unchanged": cfg["sp500_candidate"].get(
                "existing_approved_baseline_unchanged", False
            ),
        },
        "warnings": warnings,
    }


def print_status_human(payload: dict[str, Any]) -> None:
    universe = payload["universe"]
    print("Data update status")
    print(f"- production writes: {'enabled' if payload['production_writes_enabled'] else 'disabled (shadow only)'}")
    print(f"- SPY snapshot: {universe['as_of_date']} / {universe['constituents']} equities")
    approved_floor = universe["earliest_approved_date"]
    approved_latest = universe["latest_approved_date"]
    approved_range = approved_latest if approved_floor == approved_latest else f"{approved_floor} to {approved_latest}"
    print(
        f"- approved current-member price dates: {approved_range} "
        f"(latest completed session: {universe['latest_completed_session']})"
    )
    print(f"- approved current-member coverage: {universe['price_available']}/{universe['constituents']}")
    tracked = payload["tracked_assets"]
    tracked_text = ", ".join(
        f"{item['symbol']}={item.get('latest_date') or 'missing'} ({item.get('registered_status', 'unregistered')})"
        for item in tracked["assets"]
    )
    print(
        f"- tracked research assets: {tracked_text}; "
        f"shadow updates={'enabled' if tracked['shadow_update_enabled'] else 'not implemented'}"
    )
    capabilities = payload.get("capabilities", {})
    unavailable = [name for name, status in capabilities.items() if status != "ready"]
    print(f"- unavailable/planned capabilities: {', '.join(unavailable) if unavailable else 'none'}")
    source = payload.get("latest_source_validation") or {}
    source_summary = source.get("summary", {})
    source_sample = source_summary.get("sample_symbols")
    source_at_least_two = source_summary.get("symbols_with_at_least_two_sources")
    source_all_three = source_summary.get("symbols_with_all_three_sources")
    source_detail = (
        f"; >=2 sources {source_at_least_two}/{source_sample}; 3 sources {source_all_three}/{source_sample}"
        if source_sample is not None
        else ""
    )
    print(f"- latest source validation: {source.get('run_id', 'none')}{source_detail}")
    shadow = payload.get("latest_shadow_update") or {}
    shadow_scope = shadow.get("scope", {})
    shadow_detail = (
        f"; through {shadow.get('desired_end_session')}; {shadow_scope.get('kind')} "
        f"{shadow_scope.get('selected_count')}/{shadow_scope.get('current_universe_count')}"
        if shadow
        else ""
    )
    print(f"- latest shadow update: {shadow.get('run_id', 'none')}{shadow_detail}")
    campaign = payload["shadow_validation_campaign"]
    print(
        f"- five-session campaign: {campaign['current_consecutive_sessions']}/"
        f"{campaign['target_consecutive_completed_sessions']} consecutive XNYS sessions; "
        f"remaining={campaign['remaining_sessions']}"
    )
    schedule = payload["schedule"]
    print(
        f"- schedule: {'installed' if schedule['installed'] else 'not installed'}; "
        f"target {schedule['target_local_time']} {schedule['timezone']}"
    )
    imports = payload["purchased_imports"]
    print(f"- purchased imports: {imports['batches']} batches; statuses={imports['status_counts']}")
    nasdaq100 = payload["nasdaq100_candidate"]
    print(
        f"- Nasdaq-100 point-in-time candidate: {nasdaq100['status']}; "
        f"approved={str(nasdaq100['approved']).lower()}; build_id={nasdaq100['build_id'] or 'none'}"
    )
    sp500_candidate = payload["sp500_candidate"]
    print(
        f"- S&P 500 point-in-time candidate reaudit: {sp500_candidate['status']}; "
        f"approved={str(sp500_candidate['approved']).lower()}; "
        f"build_id={sp500_candidate['build_id'] or 'none'}; approved baseline unchanged=true"
    )
    credentials = payload["credentials"]
    print("- credentials: " + ", ".join(f"{name}={'present' if present else 'missing'}" for name, present in credentials.items()))
    if payload["warnings"]:
        print("Warnings:")
        for warning in payload["warnings"]:
            print(f"- {warning}")


def quick_check(*, verify_hashes: bool = False) -> dict[str, Any]:
    cfg = config()
    checks: list[dict[str, str]] = []

    def record(name: str, status: str, detail: str) -> None:
        checks.append({"name": name, "status": status, "detail": detail})

    try:
        rows = read_current_constituents(cfg)
        record("current_constituents", "pass", f"{len(rows)} unique equities")
    except Exception as exc:
        rows = []
        record("current_constituents", "fail", str(exc))
    if rows:
        try:
            freshness = database_freshness(rows, cfg)
            missing_expected = [row["symbol"] for row in rows if row["vendor_price_available"].lower() == "false"]
            unexpected_missing = sorted(set(freshness["price_missing"]) - set(missing_expected))
            if unexpected_missing or freshness["symbol_mismatches"]:
                record(
                    "approved_current_member_prices",
                    "fail",
                    f"unexpected missing={unexpected_missing}; symbol mismatches={freshness['symbol_mismatches']}",
                )
            else:
                record(
                    "approved_current_member_prices",
                    "pass",
                    f"{freshness['price_available']} available; expected unavailable={missing_expected}",
                )
            approved_dates = sorted(freshness["last_date_distribution"])
            if not approved_dates:
                record("approved_freshness", "fail", "no readable approved current-member price dates")
            else:
                earliest = approved_dates[0]
                latest = approved_dates[-1]
                completed = completed_session_on_or_before().isoformat()
                status = "warning" if earliest < completed or len(approved_dates) > 1 else "pass"
                record(
                    "approved_freshness",
                    status,
                    f"floor={earliest}; latest={latest}; latest completed session={completed}; "
                    f"distinct end dates={len(approved_dates)}",
                )
        except Exception as exc:
            record("approved_current_member_prices", "fail", str(exc))
    try:
        tracked = tracked_asset_freshness(cfg)
        required = set(tracked["required_for_current_research"])
        by_symbol = {item["symbol"]: item for item in tracked["assets"]}
        missing = sorted(symbol for symbol in required if not by_symbol.get(symbol, {}).get("latest_date"))
        unapproved = sorted(
            symbol
            for symbol in required
            if by_symbol.get(symbol, {}).get("registered_status") != "approved"
        )
        if missing or unapproved:
            record("tracked_research_assets", "fail", f"missing={missing}; unapproved={unapproved}")
        else:
            dates = {symbol: by_symbol[symbol]["latest_date"] for symbol in sorted(required)}
            latest_completed = completed_session_on_or_before().isoformat()
            status = "warning" if any(value < latest_completed for value in dates.values()) else "pass"
            record(
                "tracked_research_assets",
                status,
                f"dates={dates}; shadow update enabled={tracked['shadow_update_enabled']}",
            )
    except Exception as exc:
        record("tracked_research_assets", "fail", str(exc))
    try:
        manifest = load_json(BASELINE_MANIFEST_PATH)
        quality = load_json(BASELINE_QUALITY_PATH)
        if manifest.get("baseline_status") != "validated_for_point_in_time_research_with_documented_gaps":
            raise UpdateError(f"unexpected baseline status: {manifest.get('baseline_status')}")
        blocking = quality.get("summary", {}).get("blocking_issue_counts", {})
        if any(blocking.values()):
            raise UpdateError(f"blocking baseline issues: {blocking}")
        record("baseline_manifest", "pass", f"build_id={manifest.get('build_id')}")
    except Exception as exc:
        record("baseline_manifest", "fail", str(exc))
    try:
        identities = security_identity_reviews(cfg)
        actions = corporate_action_reviews(cfg)
        record(
            "operational_review_registries",
            "pass",
            f"identity reviews={len(identities)}; corporate-action reviews={len(actions)}",
        )
    except Exception as exc:
        record("operational_review_registries", "fail", str(exc))
    try:
        nasdaq_registry = load_json(NASDAQ100_REGISTRY_PATH)
        if nasdaq_registry.get("review_status") != "pending_review":
            raise UpdateError("Nasdaq-100 registry is not pending_review")
        if nasdaq_registry.get("production_writes_enabled") is not False:
            raise UpdateError("Nasdaq-100 registry permits production writes")
        if nasdaq_registry.get("promotion_policy", {}).get("auto_promote_to_approved") is not False:
            raise UpdateError("Nasdaq-100 registry permits automatic approval")
        manifest_path = WORKSPACE_ROOT / cfg["nasdaq100_candidate"]["manifest"]
        if not manifest_path.is_file():
            record("nasdaq100_candidate", "warning", "candidate has not been built")
        else:
            manifest = load_json(manifest_path)
            if manifest.get("status") != "candidate_pending_review" or manifest.get("approved") is not False:
                raise UpdateError("Nasdaq-100 candidate status is unsafe")
            output_failures = []
            for output in (
                *manifest.get("outputs", {}).values(),
                *manifest.get("quality_outputs", {}).values(),
            ):
                path = WORKSPACE_ROOT / output["path"]
                if not path.is_file() or sha256_file(path) != output["sha256"]:
                    output_failures.append(output["path"])
            if output_failures:
                raise UpdateError(f"Nasdaq-100 candidate output hash failures: {output_failures}")
            record("nasdaq100_candidate", "pass", f"build_id={manifest.get('build_id')}; approved=false")
    except Exception as exc:
        record("nasdaq100_candidate", "fail", str(exc))
    try:
        sp500_candidate_registry = load_json(SP500_CANDIDATE_REGISTRY_PATH)
        if sp500_candidate_registry.get("review_status") != "pending_review":
            raise UpdateError("S&P 500 candidate registry is not pending_review")
        if sp500_candidate_registry.get("production_writes_enabled") is not False:
            raise UpdateError("S&P 500 candidate registry permits production writes")
        if sp500_candidate_registry.get("promotion_policy", {}).get("auto_promote_to_approved") is not False:
            raise UpdateError("S&P 500 candidate registry permits automatic approval")
        if cfg["sp500_candidate"].get("production_price_output") is not None:
            raise UpdateError("S&P 500 candidate is configured with a production price output")
        manifest_path = WORKSPACE_ROOT / cfg["sp500_candidate"]["manifest"]
        if not manifest_path.is_file():
            record("sp500_candidate", "warning", "candidate reaudit has not been built")
        else:
            manifest = load_json(manifest_path)
            if manifest.get("status") != "candidate_pending_review" or manifest.get("approved") is not False:
                raise UpdateError("S&P 500 candidate status is unsafe")
            output_failures = []
            for output in (
                *manifest.get("outputs", {}).values(),
                *manifest.get("quality_outputs", {}).values(),
            ):
                path = WORKSPACE_ROOT / output["path"]
                if not path.is_file() or sha256_file(path) != output["sha256"]:
                    output_failures.append(output["path"])
            if output_failures:
                raise UpdateError(f"S&P 500 candidate output hash failures: {output_failures}")
            record(
                "sp500_candidate",
                "pass",
                f"build_id={manifest.get('build_id')}; approved=false; existing baseline unchanged",
            )
    except Exception as exc:
        record("sp500_candidate", "fail", str(exc))
    source_value, source_issues = verify_pointer(SOURCE_VALIDATION_LATEST)
    if source_issues:
        record("source_validation_pointer", "fail", "; ".join(source_issues))
    else:
        record("source_validation_pointer", "pass", f"run_id={source_value.get('run_id') if source_value else None}")
        source_summary = (source_value or {}).get("summary", {})
        sample = int(source_summary.get("sample_symbols", 0))
        all_three = int(source_summary.get("symbols_with_all_three_sources", 0))
        at_least_two = int(source_summary.get("symbols_with_at_least_two_sources", 0))
        coverage_status = "pass" if sample and all_three == sample else "warning"
        record(
            "source_validation_coverage",
            coverage_status,
            f"at least two sources={at_least_two}/{sample}; all three sources={all_three}/{sample}",
        )

    shadow_pointer = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"] / "latest.json"
    shadow_value, shadow_issues = verify_pointer(shadow_pointer)
    if not shadow_pointer.exists():
        record("shadow_update_pointer", "warning", "no shadow update has completed yet")
    elif shadow_issues:
        record("shadow_update_pointer", "fail", "; ".join(shadow_issues))
    else:
        shadow_value = shadow_value or {}
        record("shadow_update_pointer", "pass", f"run_id={shadow_value.get('run_id')}")
        scope = shadow_value.get("scope", {})
        scope_kind = scope.get("kind", "unknown")
        selected_count = scope.get("selected_count")
        current_count = scope.get("current_universe_count")
        scope_status = "pass" if scope_kind == "full_current_universe" else "warning"
        record(
            "shadow_update_scope",
            scope_status,
            f"kind={scope_kind}; selected={selected_count}; current universe={current_count}",
        )
        shadow_summary = shadow_value.get("summary", {})
        production_rows = int(shadow_summary.get("production_rows_written", 0))
        record(
            "shadow_write_boundary",
            "pass" if production_rows == 0 else "fail",
            f"production rows written={production_rows}; production writes enabled={cfg['production_writes_enabled']}",
        )
        campaign = shadow_campaign_status(cfg)
        record(
            "shadow_validation_campaign",
            "pass" if campaign["complete"] else "warning",
            f"consecutive XNYS sessions={campaign['current_consecutive_sessions']}/"
            f"{campaign['target_consecutive_completed_sessions']}; remaining={campaign['remaining_sessions']}",
        )
    for name in KEYCHAIN_SERVICES:
        record(f"credential_{name}", "pass" if credential(name) else "warning", "present" if credential(name) else "missing")
    if verify_hashes:
        try:
            registry = load_json(BASELINE_REGISTRY_PATH)
            targets = [registry["price_source"], registry["current_holdings_source"], *registry["frozen_existing_outputs"]]
            for item in targets:
                path = WORKSPACE_ROOT / item["path"]
                if sha256_file(path) != item["sha256"]:
                    raise UpdateError(f"SHA256 mismatch: {item['path']}")
            record("frozen_source_hashes", "pass", f"verified {len(targets)} frozen files")
        except Exception as exc:
            record("frozen_source_hashes", "fail", str(exc))
        try:
            manifest = load_json(BASELINE_MANIFEST_PATH)
            canonical_files = manifest.get("price_files", [])
            expected_count = manifest.get("summary", {}).get("instrument_count")
            if not canonical_files or len(canonical_files) != expected_count:
                raise UpdateError(
                    f"canonical price manifest count mismatch: files={len(canonical_files)}; expected={expected_count}"
                )
            for item in canonical_files:
                path = WORKSPACE_ROOT / item["path"]
                if not path.is_file() or sha256_file(path) != item["canonical_sha256"]:
                    raise UpdateError(f"Canonical price SHA256 mismatch: {item['path']}")
            record("approved_price_hashes", "pass", f"verified {len(canonical_files)} approved price files")
        except Exception as exc:
            record("approved_price_hashes", "fail", str(exc))
        try:
            imported = load_json(WORKSPACE_ROOT / cfg["purchased_import"]["registry"])
            checked = 0
            for batch in imported.get("batches", []):
                root = WORKSPACE_ROOT / batch["destination"]
                for item in batch["files"]:
                    path = root / item["relative_path"]
                    if not path.is_file() or sha256_file(path) != item["sha256"]:
                        raise UpdateError(f"Imported file mismatch: {path}")
                    checked += 1
            record("purchased_import_hashes", "pass", f"verified {checked} imported files")
        except Exception as exc:
            record("purchased_import_hashes", "fail", str(exc))
        try:
            shadow_value, shadow_issues = verify_pointer(
                WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"] / "latest.json"
            )
            if shadow_issues or not shadow_value:
                raise UpdateError("; ".join(shadow_issues) or "No latest shadow update")
            report = load_json(WORKSPACE_ROOT / shadow_value["report_json"])
            candidates, archives = verify_shadow_report_artifacts(report)
            record(
                "shadow_artifact_hashes",
                "pass",
                f"verified {candidates} candidate files and {archives} unique source archives",
            )
        except Exception as exc:
            record("shadow_artifact_hashes", "fail", str(exc))
    summary = Counter(item["status"] for item in checks)
    return {
        "schema_version": 1,
        "generated_at_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "mode": "deep_hash" if verify_hashes else "quick",
        "summary": dict(summary),
        "checks": checks,
        "status": "fail" if summary["fail"] else ("warning" if summary["warning"] else "pass"),
    }


def run_deep_process_checks() -> list[dict[str, Any]]:
    commands = (
        ("raw_inventory", [str(BACKTEST_ROOT / ".venv/bin/python"), "scripts/inventory_raw_data.py", "--verify"]),
        ("workspace_audit", [str(BACKTEST_ROOT / ".venv/bin/python"), "-m", "scripts.audit_workspace"]),
        ("unit_tests", [str(BACKTEST_ROOT / ".venv/bin/python"), "-m", "unittest", "discover", "-s", "tests"]),
    )
    results = []
    for name, command in commands:
        started = time.monotonic()
        result = subprocess.run(command, cwd=BACKTEST_ROOT, capture_output=True, text=True)
        results.append(
            {
                "name": name,
                "status": "pass" if result.returncode == 0 else "fail",
                "returncode": result.returncode,
                "elapsed_seconds": round(time.monotonic() - started, 3),
                "output_tail": "\n".join((result.stdout + result.stderr).strip().splitlines()[-12:]),
            }
        )
    return results


def print_check_human(payload: dict[str, Any]) -> None:
    print(f"Data integrity check: {payload['status'].upper()} ({payload['mode']})")
    for item in payload["checks"]:
        print(f"- [{item['status']}] {item['name']}: {item['detail']}")
    for item in payload.get("process_checks", []):
        print(f"- [{item['status']}] {item['name']} ({item['elapsed_seconds']}s)")
        if item["status"] == "fail":
            print(item["output_tail"])


class TwelveMinuteLimiter:
    def __init__(self, credits_per_minute: int = 8) -> None:
        self.limit = credits_per_minute
        self.marker: str | None = None
        self.used = 0

    @staticmethod
    def current_marker() -> str:
        return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%MZ")

    def refresh(self) -> None:
        marker = self.current_marker()
        if marker != self.marker:
            self.marker = marker
            self.used = 0

    def capacity(self) -> int:
        self.refresh()
        return self.limit - self.used

    def consume(self, credits: int) -> None:
        self.refresh()
        self.used += credits

    def wait_next_minute(self) -> None:
        marker = self.current_marker()
        while self.current_marker() == marker:
            time.sleep(5)
        self.marker = self.current_marker()
        self.used = 0


def fetch_twelve_adjusted(
    symbols: list[str],
    start: str,
    end: str,
    api_key: str,
    raw_run: Path,
    limiter: TwelveMinuteLimiter,
    archive_prefix: str,
) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]], list[dict[str, str]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    batches: list[dict[str, Any]] = []
    symbol_failures: list[dict[str, str]] = []
    offset = 0
    batch_number = 0
    rate_waits = 0
    while offset < len(symbols):
        if limiter.capacity() <= 0:
            limiter.wait_next_minute()
        batch = symbols[offset : offset + limiter.capacity()]
        payload, call = curl_json(
            TWELVE_URL,
            params={
                "symbol": ",".join(batch),
                "interval": "1day",
                "start_date": start,
                "end_date": end,
                "outputsize": "5000",
                "adjust": "all",
            },
            headers={"Authorization": f"apikey {api_key}"},
        )
        message = str(payload.get("message", "")) if isinstance(payload, dict) else ""
        if isinstance(payload, dict) and payload.get("status") == "error" and "current minute" in message.lower():
            rate_waits += 1
            if rate_waits > 3:
                raise ProviderError(f"Twelve Data minute limit did not recover: {message}")
            limiter.wait_next_minute()
            continue
        if isinstance(payload, dict) and payload.get("status") == "error":
            raise ProviderError(f"Twelve Data error: {message or 'unknown error'}")
        batch_number += 1
        archive = archive_payload(raw_run, "twelve_data", f"{archive_prefix}_batch_{batch_number:03d}", payload)
        for symbol in batch:
            try:
                rows = parse_twelve(payload, symbol)
                validate_rows(rows, "twelve_adjusted", symbol)
                result[symbol] = rows
            except ProviderError as exc:
                symbol_failures.append({"symbol": symbol, "error": str(exc)})
        batches.append({"symbols": batch, "credits": len(batch), "call": call, "archive": archive})
        limiter.consume(len(batch))
        offset += len(batch)
    return result, batches, symbol_failures


def load_twelve_archive_data(
    replay_run_id: str,
    selected_symbols: list[str],
) -> tuple[
    dict[str, list[dict[str, Any]]],
    list[dict[str, Any]],
    list[dict[str, str]],
    list[str],
]:
    if not re.fullmatch(r"run_[A-Za-z0-9_.-]+", replay_run_id):
        raise UpdateError("Twelve Data archive source must be an existing immutable run_* ID")
    raw_root = WORKSPACE_ROOT / config()["shadow_update"]["raw_archive_output"] / replay_run_id / "twelve_data"
    if not raw_root.is_dir():
        raise UpdateError(f"Twelve Data replay archive is missing: {raw_root.relative_to(WORKSPACE_ROOT)}")
    remaining = set(selected_symbols)
    result: dict[str, list[dict[str, Any]]] = {}
    batches: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    for path in sorted(raw_root.glob("*.json")):
        payload = load_json(path)
        found: list[str] = []
        for symbol in sorted(remaining):
            if symbol not in payload and payload.get("meta", {}).get("symbol") != symbol:
                continue
            try:
                rows = parse_twelve(payload, symbol)
                validate_rows(rows, "twelve_adjusted_replay", symbol)
                result[symbol] = rows
                found.append(symbol)
            except ProviderError as exc:
                failures.append({"symbol": symbol, "error": str(exc)})
        if found:
            remaining.difference_update(found)
            batches.append(
                {
                    "symbols": found,
                    "credits": 0,
                    "replay": True,
                    "source_archive": {
                        "path": str(path.relative_to(WORKSPACE_ROOT)),
                        "sha256": sha256_file(path),
                        "bytes": path.stat().st_size,
                    },
                }
            )
    return result, batches, failures, sorted(remaining)


def prepare_twelve_resume_data(
    archived_data: dict[str, list[dict[str, Any]]],
    selected_symbols: list[str],
    desired_end: date,
    refresh_symbols: Iterable[str] = (),
) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, str]], list[str]]:
    """Select reusable archived series without mutating the immutable source archive."""
    reusable = dict(archived_data)
    issues: list[dict[str, str]] = []
    stale = [
        symbol
        for symbol, rows in reusable.items()
        if not rows or date.fromisoformat(rows[-1]["date"]) < desired_end
    ]
    for symbol in stale:
        last_date = reusable[symbol][-1]["date"] if reusable[symbol] else None
        issues.append(
            {
                "symbol": symbol,
                "error": f"Archived series is stale for {desired_end}: last_date={last_date}",
            }
        )
        reusable.pop(symbol, None)
    forced = set(refresh_symbols)
    for symbol in forced:
        reusable.pop(symbol, None)
    missing = sorted(set(selected_symbols) - set(reusable) | forced)
    return reusable, issues, missing


def load_twelve_replay(
    replay_run_id: str,
    selected_symbols: list[str],
) -> tuple[dict[str, list[dict[str, Any]]], list[dict[str, Any]], list[dict[str, str]]]:
    result, batches, failures, missing = load_twelve_archive_data(replay_run_id, selected_symbols)
    failures.extend(
        {"symbol": symbol, "error": f"No replay payload found in {replay_run_id}"}
        for symbol in missing
    )
    return result, batches, failures


def select_symbols(requested: str | None, limit: int | None, current: list[dict[str, str]]) -> list[dict[str, str]]:
    by_symbol = {row["symbol"]: row for row in current}
    if requested:
        symbols = [value.strip().upper() for value in requested.split(",") if value.strip()]
        missing = sorted(set(symbols) - set(by_symbol))
        if missing:
            raise UpdateError(f"Not in current dated SPY snapshot: {missing}")
        selected = [by_symbol[symbol] for symbol in symbols]
    else:
        selected = current
    if limit is not None:
        if limit <= 0:
            raise UpdateError("--limit must be positive")
        selected = selected[:limit]
    return selected


def choose_tiingo_rotation(
    selected_symbols: list[str],
    cfg: dict[str, Any],
    run_date: date,
    mode: str,
    priority_symbols: Iterable[str] = (),
) -> tuple[list[str], dict[str, Any]]:
    if mode == "none":
        return [], {"mode": mode, "monthly_unique_before": 0, "monthly_unique_after": 0}
    if mode == "all":
        return selected_symbols, {"mode": mode, "monthly_unique_before": None, "monthly_unique_after": None}
    root = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"]
    month_prefix = run_date.strftime("%Y%m")
    month_serial = run_date.year * 12 + run_date.month - 1
    monthly_budget = int(cfg["shadow_update"]["tiingo_monthly_unique_budget"])
    rotation_offset = (month_serial * monthly_budget) % len(selected_symbols) if selected_symbols else 0
    rotated_symbols = selected_symbols[rotation_offset:] + selected_symbols[:rotation_offset]
    used: set[str] = set()
    if root.is_dir():
        for report_path in root.glob(f"run_{month_prefix}*/report.json"):
            try:
                prior = load_json(report_path)
                used.update(prior.get("tiingo", {}).get("requested_symbols", []))
            except (OSError, json.JSONDecodeError):
                continue
    budget = monthly_budget
    batch_size = int(cfg["shadow_update"]["tiingo_daily_batch_size"])
    mandatory = [symbol for symbol in cfg["shadow_update"]["tiingo_mandatory_symbols"] if symbol in selected_symbols]
    priorities: list[str] = []
    for symbol in [*mandatory, *priority_symbols]:
        if symbol in selected_symbols and symbol not in priorities:
            priorities.append(symbol)
    remaining_new_budget = max(0, budget - len(used))
    new_candidates = [symbol for symbol in rotated_symbols if symbol not in used]
    repeat_candidates = [symbol for symbol in rotated_symbols if symbol in used]
    chosen: list[str] = []
    for symbol in priorities + new_candidates[:remaining_new_budget] + repeat_candidates:
        if symbol not in chosen:
            chosen.append(symbol)
        if len(chosen) >= batch_size:
            break
    unique_after = used | set(chosen)
    if len(unique_after) > budget:
        overflow = len(unique_after) - budget
        removable = [symbol for symbol in reversed(chosen) if symbol not in used]
        for symbol in removable[:overflow]:
            chosen.remove(symbol)
        unique_after = used | set(chosen)
    return chosen, {
        "mode": mode,
        "monthly_unique_budget": budget,
        "monthly_rotation_offset": rotation_offset,
        "monthly_unique_before": len(used),
        "monthly_unique_after": len(unique_after),
        "priority_symbols": priorities,
    }


def replay_run_path(run_id: str, cfg: dict[str, Any], provider: str | None = None) -> Path:
    if not re.fullmatch(r"run_[A-Za-z0-9_.-]+", run_id):
        raise UpdateError("Replay run IDs must be immutable run_* IDs")
    if provider:
        path = WORKSPACE_ROOT / cfg["shadow_update"]["raw_archive_output"] / run_id / provider
    else:
        path = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"] / run_id / "report.json"
    if not path.exists():
        raise UpdateError(f"Replay source is missing: {path.relative_to(WORKSPACE_ROOT)}")
    return path


def choose_tiingo_replay(
    selected_symbols: list[str],
    cfg: dict[str, Any],
    replay_run_ids: list[str],
    priority_symbols: Iterable[str] = (),
) -> tuple[list[str], dict[str, Any]]:
    requested_from_sources: list[str] = []
    for run_id in replay_run_ids:
        report = load_json(replay_run_path(run_id, cfg))
        for symbol in report.get("tiingo", {}).get("requested_symbols", []):
            if symbol in selected_symbols and symbol not in requested_from_sources:
                requested_from_sources.append(symbol)
    mandatory = [symbol for symbol in cfg["shadow_update"]["tiingo_mandatory_symbols"] if symbol in selected_symbols]
    ordered: list[str] = []
    for symbol in [*mandatory, *priority_symbols, *requested_from_sources, *selected_symbols]:
        if symbol in selected_symbols and symbol not in ordered:
            ordered.append(symbol)
    batch_size = int(cfg["shadow_update"]["tiingo_daily_batch_size"])
    chosen = ordered[:batch_size]
    return chosen, {
        "mode": "replay_then_fetch_missing",
        "source_run_ids": replay_run_ids,
        "priority_symbols": [symbol for symbol in ordered if symbol in set(mandatory) | set(priority_symbols)],
        "requested_from_source_reports": len(requested_from_sources),
        "selected_count": len(chosen),
    }


def load_tiingo_replay_payload(
    symbol: str,
    cfg: dict[str, Any],
    replay_run_ids: list[str],
) -> tuple[Any, dict[str, Any]] | None:
    for run_id in reversed(replay_run_ids):
        provider_root = replay_run_path(run_id, cfg, "tiingo")
        path = provider_root / f"{symbol}.json"
        if not path.is_file():
            continue
        return load_json(path), {
            "replay": True,
            "source_run_id": run_id,
            "source_archive": {
                "path": str(path.relative_to(WORKSPACE_ROOT)),
                "sha256": sha256_file(path),
                "bytes": path.stat().st_size,
            },
        }
    return None


def format_number(value: float) -> str:
    return format(value, ".12g")


def stable_return_tail(
    common: list[str],
    approved: dict[str, dict[str, Any]],
    provider: dict[str, dict[str, Any]],
    tolerance: float,
) -> tuple[list[str], list[dict[str, Any]]]:
    comparisons: list[dict[str, Any]] = []
    for prior, current in zip(common, common[1:]):
        approved_return = approved[current]["close"] / approved[prior]["close"] - 1
        provider_return = provider[current]["close"] / provider[prior]["close"] - 1
        comparisons.append(
            {
                "prior_date": prior,
                "date": current,
                "abs_diff": abs(approved_return - provider_return),
            }
        )
    tail_start = len(comparisons)
    while tail_start > 0 and comparisons[tail_start - 1]["abs_diff"] <= tolerance:
        tail_start -= 1
    tail_comparisons = comparisons[tail_start:]
    tail_dates = [tail_comparisons[0]["prior_date"], *[item["date"] for item in tail_comparisons]] if tail_comparisons else []
    historical_outliers = [item for item in comparisons[:tail_start] if item["abs_diff"] > tolerance]
    return tail_dates, historical_outliers


def cross_check_window_start(
    analysis: dict[str, Any],
    desired_end: date,
    settings: dict[str, Any],
) -> date:
    start = desired_end - timedelta(days=int(settings.get("cross_check_calendar_days", 60)))
    candidate_start = analysis.get("candidate_start_date")
    if candidate_start:
        start = max(start, date.fromisoformat(candidate_start))
    return start


def evaluate_cross_check_coverage(
    analysis: dict[str, Any],
    expected_cross_dates: list[str],
    common_dates: list[str],
    stable_tail_sessions: int,
    latest_common_date: str | None,
    desired_end: date,
    settings: dict[str, Any],
) -> tuple[bool, int, int, bool]:
    if analysis.get("kind") == "bootstrap":
        required_common_dates = len(expected_cross_dates)
        required_stable_tail_sessions = required_common_dates
        complete_cross_coverage = common_dates == expected_cross_dates
        passed = (
            required_common_dates >= 2
            and complete_cross_coverage
            and stable_tail_sessions >= required_stable_tail_sessions
            and latest_common_date == desired_end.isoformat()
        )
    else:
        required_common_dates = int(settings["minimum_overlap_sessions"])
        required_stable_tail_sessions = int(settings["minimum_stable_tail_sessions"])
        complete_cross_coverage = True
        passed = (
            len(common_dates) >= required_common_dates
            and stable_tail_sessions >= required_stable_tail_sessions
            and latest_common_date == desired_end.isoformat()
        )
    return passed, required_common_dates, required_stable_tail_sessions, complete_cross_coverage


def analyze_candidate(
    row: dict[str, str],
    provider_rows: list[dict[str, Any]],
    desired_end: date,
    cfg: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    settings = cfg["shadow_update"]
    symbol = row["symbol"]
    price_root = WORKSPACE_ROOT / cfg["universe"]["approved_price_directory"]
    instrument = row["vendor_instrument_id"]
    available = row["vendor_price_available"].lower() == "true" and bool(instrument)
    provider_rows = [item for item in provider_rows if date.fromisoformat(item["date"]) <= desired_end]
    if not provider_rows:
        raise ProviderError(f"Twelve Data returned no rows through {desired_end} for {symbol}")
    latest_provider = date.fromisoformat(provider_rows[-1]["date"])
    if not available:
        provider_first = date.fromisoformat(provider_rows[0]["date"])
        identity_review = security_identity_reviews(cfg).get(symbol)
        candidate_start = max(date.fromisoformat(settings["bootstrap_start"]), provider_first)
        if identity_review and identity_review.get("decision") == "verified_for_shadow_bootstrap_only":
            candidate_start = max(candidate_start, date.fromisoformat(identity_review["candidate_not_before"]))
        expected_tail = expected_sessions(candidate_start, desired_end)
        expected_set = set(expected_tail)
        candidate_provider_dates = {
            item["date"] for item in provider_rows if candidate_start <= date.fromisoformat(item["date"]) <= desired_end
        }
        missing = [day for day in expected_tail if day not in candidate_provider_dates]
        unexpected = sorted(candidate_provider_dates - expected_set)
        identity_verified = bool(identity_review and identity_review.get("decision") == "verified_for_shadow_bootstrap_only")
        if identity_verified and latest_provider >= desired_end and not missing and not unexpected:
            status = "ready_bootstrap_candidate"
        elif latest_provider >= desired_end and not missing and not unexpected:
            status = "bootstrap_needs_identity_review"
        else:
            status = "review"
        candidate_rows = [
            {
                "date": item["date"],
                "symbol": symbol,
                "open": format_number(item["open"]),
                "high": format_number(item["high"]),
                "low": format_number(item["low"]),
                "close": format_number(item["close"]),
                "volume": format_number(item["volume"]),
            }
            for item in provider_rows
            if item["date"] in expected_set
        ]
        return (
            {
                "symbol": symbol,
                "kind": "bootstrap",
                "status": status,
                "approved_last_date": None,
                "provider_first_date": provider_rows[0]["date"],
                "provider_last_date": provider_rows[-1]["date"],
                "anchor_date": None,
                "scale_factor": None,
                "overlap_sessions": 0,
                "overlap_return_abs_diff_p95": None,
                "overlap_return_abs_diff_max": None,
                "expected_new_sessions": len(expected_tail),
                "candidate_rows": len(candidate_rows),
                "candidate_start_date": candidate_start.isoformat(),
                "missing_expected_dates": missing,
                "unexpected_dates": unexpected,
                "identity_review_status": identity_review.get("decision") if identity_review else "missing",
                "production_approved": bool(identity_review and identity_review.get("production_approved")),
                "cross_check_status": "not_selected",
            },
            candidate_rows,
        )
    path = price_root / f"{instrument}.csv"
    last = read_last_csv_record(path)
    last_date = date.fromisoformat(last["date"])
    overlap_start = last_date - timedelta(days=int(settings["overlap_calendar_days"]))
    approved_rows = read_canonical_from(path, overlap_start)
    approved = {item["date"]: item for item in approved_rows}
    provider = {item["date"]: item for item in provider_rows}
    common = sorted(set(approved) & set(provider))
    missing_overlap = sorted(set(approved) - set(provider))
    unexpected_overlap = sorted(
        day for day in set(provider) - set(approved) if overlap_start <= date.fromisoformat(day) <= last_date
    )
    return_diffs = [
        abs(
            approved[current]["close"] / approved[prior]["close"]
            - provider[current]["close"] / provider[prior]["close"]
        )
        for prior, current in zip(common, common[1:])
    ]
    anchor = common[-1] if common else None
    scale = approved[anchor]["close"] / provider[anchor]["close"] if anchor else None
    new_expected = expected_sessions(last_date + timedelta(days=1), desired_end)
    new_expected_set = set(new_expected)
    new_provider = [item for item in provider_rows if item["date"] in new_expected_set]
    new_dates = {item["date"] for item in new_provider}
    missing = [day for day in new_expected if day not in new_dates]
    unexpected_new = sorted(
        item["date"]
        for item in provider_rows
        if last_date < date.fromisoformat(item["date"]) <= desired_end and item["date"] not in new_expected_set
    )
    p95 = percentile(return_diffs, 0.95)
    maximum = max(return_diffs) if return_diffs else None
    tolerance = float(settings["overlap_close_return_abs_diff_tolerance"])
    stable_tail_dates, historical_outliers = stable_return_tail(common, approved, provider, tolerance)
    action_review = corporate_action_reviews(cfg).get(symbol) if historical_outliers else None
    documented_outlier_dates = set(action_review.get("covered_return_outlier_dates", [])) if action_review else set()
    undocumented_outliers = [
        item for item in historical_outliers if item["date"] not in documented_outlier_dates
    ]
    stable_tail_minimum = int(settings.get("minimum_stable_tail_sessions", settings["minimum_overlap_sessions"]))
    reasons: list[str] = []
    if len(common) < int(settings["minimum_overlap_sessions"]):
        reasons.append("insufficient_overlap")
    if missing_overlap or unexpected_overlap:
        reasons.append("overlap_date_mismatch")
    if anchor != last["date"]:
        reasons.append("anchor_not_approved_tail")
    if len(stable_tail_dates) < stable_tail_minimum:
        reasons.append("insufficient_stable_return_tail")
    if undocumented_outliers:
        reasons.append("undocumented_historical_adjustment_basis_change")
    if missing:
        reasons.append("missing_expected_dates")
    if unexpected_new:
        reasons.append("unexpected_new_dates")
    if latest_provider < desired_end:
        reasons.append("provider_stale")
    status = "ready_candidate" if not reasons else "review"
    candidate_rows = []
    if scale is not None:
        for item in new_provider:
            candidate_rows.append(
                {
                    "date": item["date"],
                    "symbol": instrument,
                    "open": format_number(item["open"] * scale),
                    "high": format_number(item["high"] * scale),
                    "low": format_number(item["low"] * scale),
                    "close": format_number(item["close"] * scale),
                    "volume": format_number(item["volume"]),
                }
            )
    return (
        {
            "symbol": symbol,
            "kind": "incremental",
            "status": status,
            "review_reasons": reasons,
            "approved_last_date": last["date"],
            "provider_first_date": provider_rows[0]["date"],
            "provider_last_date": provider_rows[-1]["date"],
            "anchor_date": anchor,
            "scale_factor": scale,
            "overlap_sessions": len(common),
            "missing_overlap_dates": missing_overlap,
            "unexpected_overlap_dates": unexpected_overlap,
            "overlap_return_abs_diff_p95": p95,
            "overlap_return_abs_diff_max": maximum,
            "stable_tail_sessions": len(stable_tail_dates),
            "stable_tail_start_date": stable_tail_dates[0] if stable_tail_dates else None,
            "stable_tail_return_abs_diff_max": (
                max(
                    abs(
                        approved[current]["close"] / approved[prior]["close"]
                        - provider[current]["close"] / provider[prior]["close"]
                    )
                    for prior, current in zip(stable_tail_dates, stable_tail_dates[1:])
                )
                if len(stable_tail_dates) >= 2
                else None
            ),
            "historical_overlap_outliers": historical_outliers,
            "undocumented_historical_overlap_outliers": undocumented_outliers,
            "corporate_action_review_status": (
                action_review.get("decision") if action_review else "not_registered"
            ),
            "warnings": ["historical_adjustment_basis_change"] if historical_outliers else [],
            "expected_new_sessions": len(new_expected),
            "candidate_rows": len(candidate_rows),
            "missing_expected_dates": missing,
            "unexpected_new_dates": unexpected_new,
            "cross_check_status": "not_selected",
        },
        candidate_rows,
    )


def shadow_markdown(report: dict[str, Any]) -> str:
    summary = report["summary"]
    twelve = report["twelve_data"]
    lines = [
        "# SPY 当前成分影子增量更新",
        "",
        f"- Run：`{report['run_id']}`",
        f"- 目标完成交易日：{report['desired_end_session']}",
        f"- 范围：{summary['symbols']} 个；候选可用：{summary['ready']}；需复核：{summary['review']}；抓取失败：{summary['fetch_failed']}。",
        f"- 候选新增/引导行：{summary['candidate_rows']}；正式写入：0（`production_writes_enabled=false`）。",
        f"- Tiingo 交叉验证：{summary['cross_checked']} 个；通过：{summary['cross_check_passed']}。",
    ]
    if twelve.get("resume_source_run_id"):
        reused_symbols = twelve.get("resume_reused_symbols")
        reused_count = (
            len(reused_symbols)
            if reused_symbols is not None
            else sum(len(batch["symbols"]) for batch in twelve["batches"] if batch.get("replay"))
        )
        lines.append(
            f"- Twelve Data 断点来源：`{twelve['resume_source_run_id']}`；复用 {reused_count} 个标的；"
            f"来源问题 {len(twelve.get('resume_source_issues', []))} 个。"
        )
        if twelve.get("refresh_symbols"):
            lines.append(f"- Twelve Data 强制重抓：{', '.join(twelve['refresh_symbols'])}。")
    elif twelve.get("replay_source_run_id"):
        lines.append(f"- Twelve Data 重放来源：`{twelve['replay_source_run_id']}`。")
    lines.extend(
        [
            "",
            "## 需复核或失败",
            "",
            "| Symbol | 类型 | 状态 | Approved 截止 | Provider 截止 | 候选行 | 原因 |",
            "|---|---|---|---|---|---:|---|",
        ]
    )
    issues = [
        item
        for item in report["symbols"]
        if item["status"] not in {"ready_candidate", "ready_bootstrap_candidate"}
        or item.get("cross_check_status") in {"review", "failed"}
    ]
    if not issues:
        lines.append("| — | — | — | — | — | — | 无 |")
    for item in issues:
        reasons = list(item.get("review_reasons", []))
        if item.get("cross_check_status") in {"review", "failed"}:
            reasons.append(f"tiingo_cross_check_{item['cross_check_status']}")
        reason = ", ".join(reasons) or item.get("error", "")
        lines.append(
            f"| {item['symbol']} | {item.get('kind', '')} | {item['status']} | {item.get('approved_last_date') or ''} | "
            f"{item.get('provider_last_date') or ''} | {item.get('candidate_rows', 0)} | {reason} |"
        )
    lines.extend(
        [
            "",
            "## 边界",
            "",
            "本 run 仅保存供应商原始响应、缩放后的候选增量和质量证据，不修改批准价格文件。成交量保持供应商值且只作警告；候选提升需在连续影子运行达标后另行实现和授权。",
            "",
        ]
    )
    return "\n".join(lines)


def run_shadow_update(args: argparse.Namespace) -> tuple[dict[str, Any], int]:
    cfg = config()
    api_key = credential("TWELVE_DATA_API_KEY")
    tiingo_token = credential("TIINGO_API_TOKEN")
    if args.replay_twelve_run and args.resume_twelve_run:
        raise UpdateError("Use only one of --replay-twelve-run and --resume-twelve-run")
    if args.refresh_symbols and not args.resume_twelve_run:
        raise UpdateError("--refresh-symbols requires --resume-twelve-run")
    if not api_key and not args.replay_twelve_run:
        raise UpdateError("TWELVE_DATA_API_KEY is missing from environment and macOS Keychain")
    current = read_current_constituents(cfg)
    selected = select_symbols(args.symbols, args.limit, current)
    refresh_symbols = {
        row["symbol"] for row in select_symbols(args.refresh_symbols, None, selected)
    } if args.refresh_symbols else set()
    desired_end = completed_session_on_or_before(date.fromisoformat(args.end) if args.end else None)
    now = datetime.now(timezone.utc)
    run_id = args.run_id or now.strftime("run_%Y%m%dT%H%M%SZ")
    processed_root = WORKSPACE_ROOT / cfg["shadow_update"]["candidate_output"]
    processed_run = processed_root / run_id
    raw_run = WORKSPACE_ROOT / cfg["shadow_update"]["raw_archive_output"] / run_id
    if processed_run.exists() or raw_run.exists():
        raise UpdateError(f"Immutable run already exists: {run_id}")
    price_root = WORKSPACE_ROOT / cfg["universe"]["approved_price_directory"]
    settings = cfg["shadow_update"]
    normal: list[tuple[dict[str, str], date]] = []
    bootstrap: list[dict[str, str]] = []
    for row in selected:
        instrument = row["vendor_instrument_id"]
        if row["vendor_price_available"].lower() == "true" and instrument:
            last = read_last_csv_record(price_root / f"{instrument}.csv")
            normal.append((row, date.fromisoformat(last["date"])))
        else:
            bootstrap.append(row)
    provider_data: dict[str, list[dict[str, Any]]] = {}
    twelve_batches: list[dict[str, Any]] = []
    twelve_resume_source_issues: list[dict[str, str]] = []
    twelve_resume_reused_symbols: list[str] = []
    if args.replay_twelve_run:
        data, batches, failures = load_twelve_replay(
            args.replay_twelve_run,
            [row["symbol"] for row in selected],
        )
        provider_data.update(data)
        twelve_batches.extend(batches)
        twelve_symbol_failures = failures
    elif args.resume_twelve_run:
        selected_symbols = [row["symbol"] for row in selected]
        data, batches, source_issues, missing = load_twelve_archive_data(
            args.resume_twelve_run,
            selected_symbols,
        )
        data, resume_issues, resume_missing = prepare_twelve_resume_data(
            data,
            selected_symbols,
            desired_end,
            refresh_symbols,
        )
        source_issues.extend(resume_issues)
        missing = sorted(set(missing) | set(resume_missing))
        twelve_resume_reused_symbols = sorted(data)
        provider_data.update(data)
        twelve_batches.extend(batches)
        twelve_resume_source_issues = source_issues
        twelve_symbol_failures = []
        missing_set = set(missing)
        limiter = TwelveMinuteLimiter()
        missing_normal = [(row, last) for row, last in normal if row["symbol"] in missing_set]
        if missing_normal:
            start = min(last for _row, last in missing_normal) - timedelta(days=int(settings["overlap_calendar_days"]))
            fetched, fetched_batches, failures = fetch_twelve_adjusted(
                [row["symbol"] for row, _last in missing_normal],
                start.isoformat(),
                (desired_end + timedelta(days=1)).isoformat(),
                api_key,
                raw_run,
                limiter,
                "resume_incremental",
            )
            provider_data.update(fetched)
            twelve_batches.extend(fetched_batches)
            twelve_symbol_failures.extend(failures)
        missing_bootstrap = [row for row in bootstrap if row["symbol"] in missing_set]
        if missing_bootstrap:
            fetched, fetched_batches, failures = fetch_twelve_adjusted(
                [row["symbol"] for row in missing_bootstrap],
                settings["bootstrap_start"],
                (desired_end + timedelta(days=1)).isoformat(),
                api_key,
                raw_run,
                limiter,
                "resume_bootstrap",
            )
            provider_data.update(fetched)
            twelve_batches.extend(fetched_batches)
            twelve_symbol_failures.extend(failures)
    else:
        limiter = TwelveMinuteLimiter()
        if normal:
            start = min(last for _row, last in normal) - timedelta(days=int(settings["overlap_calendar_days"]))
            data, batches, failures = fetch_twelve_adjusted(
                [row["symbol"] for row, _last in normal],
                start.isoformat(),
                (desired_end + timedelta(days=1)).isoformat(),
                api_key,
                raw_run,
                limiter,
                "incremental",
            )
            provider_data.update(data)
            twelve_batches.extend(batches)
            twelve_symbol_failures = failures
        else:
            twelve_symbol_failures = []
        if bootstrap:
            data, batches, failures = fetch_twelve_adjusted(
                [row["symbol"] for row in bootstrap],
                settings["bootstrap_start"],
                (desired_end + timedelta(days=1)).isoformat(),
                api_key,
                raw_run,
                limiter,
                "bootstrap",
            )
            provider_data.update(data)
            twelve_batches.extend(batches)
            twelve_symbol_failures.extend(failures)
    analyses: dict[str, dict[str, Any]] = {}
    candidate_files: dict[str, dict[str, Any]] = {}
    for row in selected:
        symbol = row["symbol"]
        try:
            if symbol not in provider_data:
                detail = next(
                    (item["error"] for item in twelve_symbol_failures if item["symbol"] == symbol),
                    "Twelve Data returned no usable series",
                )
                raise ProviderError(detail)
            analysis, candidate_rows = analyze_candidate(row, provider_data[symbol], desired_end, cfg)
            analyses[symbol] = analysis
            content = csv_bytes(CANONICAL_COLUMNS, candidate_rows)
            path = processed_run / "candidates" / f"{symbol}.csv"
            atomic_write(path, content)
            candidate_files[symbol] = {
                "path": str(path.relative_to(WORKSPACE_ROOT)),
                "sha256": sha256_bytes(content),
                "rows": len(candidate_rows),
            }
        except Exception as exc:
            analyses[symbol] = {"symbol": symbol, "status": "fetch_or_analysis_failed", "error": str(exc)}
    selected_symbols = [row["symbol"] for row in selected]
    priority_cross_checks = [
        symbol
        for symbol in selected_symbols
        if analyses.get(symbol, {}).get("status") == "ready_bootstrap_candidate"
        or "historical_adjustment_basis_change" in analyses.get(symbol, {}).get("warnings", [])
    ]
    tiingo_replay_run_ids = [
        value.strip() for value in (args.replay_tiingo_runs or "").split(",") if value.strip()
    ]
    if tiingo_replay_run_ids:
        if args.tiingo == "none":
            raise UpdateError("--replay-tiingo-runs cannot be combined with --tiingo none")
        tiingo_selected, tiingo_plan = choose_tiingo_replay(
            selected_symbols,
            cfg,
            tiingo_replay_run_ids,
            priority_cross_checks,
        )
    else:
        tiingo_selected, tiingo_plan = choose_tiingo_rotation(
            selected_symbols,
            cfg,
            desired_end,
            args.tiingo,
            priority_cross_checks,
        )
    tiingo_failures: list[dict[str, str]] = []
    tiingo_metadata: dict[str, Any] = {}
    for symbol in tiingo_selected:
        try:
            twelve_rows = provider_data[symbol]
            cross_check_start = cross_check_window_start(analyses[symbol], desired_end, settings)
            cross_check_twelve = [
                item for item in twelve_rows if date.fromisoformat(item["date"]) >= cross_check_start
            ]
            start = cross_check_twelve[0]["date"]
            replayed = load_tiingo_replay_payload(symbol, cfg, tiingo_replay_run_ids) if tiingo_replay_run_ids else None
            if replayed:
                payload, call = replayed
                archive = call["source_archive"]
            else:
                if not tiingo_token:
                    raise UpdateError("TIINGO_API_TOKEN is missing and no replay payload is available")
                payload, call = curl_json(
                    TIINGO_URL.format(symbol=tiingo_symbol(symbol)),
                    params={"startDate": start, "endDate": desired_end.isoformat(), "resampleFreq": "daily"},
                    headers={"Authorization": f"Token {tiingo_token}"},
                )
                archive = archive_payload(raw_run, "tiingo", symbol, payload)
            tiingo_rows = parse_tiingo(payload, adjusted=True)
            validate_rows(tiingo_rows, "tiingo_adjusted", symbol)
            tiingo_rows = [item for item in tiingo_rows if date.fromisoformat(item["date"]) >= cross_check_start]
            metrics = compare_series(symbol, "adjusted", "twelve_data", cross_check_twelve, "tiingo", tiingo_rows)
            tolerance = float(settings["cross_source_close_return_abs_diff_tolerance"])
            twelve_by_date = {item["date"]: item for item in cross_check_twelve}
            tiingo_by_date = {item["date"]: item for item in tiingo_rows}
            cross_common = sorted(set(twelve_by_date) & set(tiingo_by_date))
            cross_stable_tail, cross_historical_outliers = stable_return_tail(
                cross_common,
                twelve_by_date,
                tiingo_by_date,
                tolerance,
            )
            metrics["stable_tail_sessions"] = len(cross_stable_tail)
            metrics["stable_tail_start_date"] = cross_stable_tail[0] if cross_stable_tail else None
            metrics["historical_return_outliers"] = cross_historical_outliers
            expected_cross_dates = expected_sessions(cross_check_start, desired_end)
            passed, required_common_dates, required_stable_tail_sessions, complete_cross_coverage = (
                evaluate_cross_check_coverage(
                    analyses[symbol],
                    expected_cross_dates,
                    cross_common,
                    len(cross_stable_tail),
                    metrics["latest_common_date"],
                    desired_end,
                    settings,
                )
            )
            metrics["required_common_dates"] = required_common_dates
            metrics["required_stable_tail_sessions"] = required_stable_tail_sessions
            metrics["complete_candidate_period_coverage"] = complete_cross_coverage
            analyses[symbol]["cross_check_status"] = "pass" if passed else "review"
            analyses[symbol]["cross_check"] = metrics
            tiingo_metadata[symbol] = {
                "requested_symbol": tiingo_symbol(symbol),
                "call": call,
                "archive": archive,
            }
        except Exception as exc:
            tiingo_failures.append({"symbol": symbol, "error": str(exc)})
            if symbol in analyses:
                analyses[symbol]["cross_check_status"] = "failed"
    symbol_reports = [analyses[row["symbol"]] for row in selected]
    ready_statuses = {"ready_candidate", "ready_bootstrap_candidate"}
    summary = {
        "symbols": len(selected),
        "ready": sum(item.get("status") in ready_statuses for item in symbol_reports),
        "review": sum(
            item.get("status") in {"review", "bootstrap_needs_cross_check", "bootstrap_needs_identity_review"}
            for item in symbol_reports
        ),
        "fetch_failed": sum(item.get("status") == "fetch_or_analysis_failed" for item in symbol_reports),
        "candidate_rows": sum(item.get("candidate_rows", 0) for item in symbol_reports),
        "cross_checked": sum(item.get("cross_check_status") in {"pass", "review", "failed"} for item in symbol_reports),
        "cross_check_passed": sum(item.get("cross_check_status") == "pass" for item in symbol_reports),
        "production_rows_written": 0,
    }
    report: dict[str, Any] = {
        "schema_version": 1,
        "run_id": run_id,
        "generated_at_utc": now.isoformat().replace("+00:00", "Z"),
        "mode": "shadow_only",
        "production_writes_enabled": False,
        "scope": {
            "kind": "full_current_universe" if len(selected) == len(current) else "partial_current_universe",
            "selected_count": len(selected),
            "current_universe_count": len(current),
            "selected_symbols": selected_symbols,
        },
        "universe_as_of_date": current[0]["as_of_date"],
        "desired_end_session": desired_end.isoformat(),
        "summary": summary,
        "settings": settings,
        "twelve_data": {
            "adjust": "all",
            "replay_source_run_id": args.replay_twelve_run,
            "resume_source_run_id": args.resume_twelve_run,
            "resume_source_issues": twelve_resume_source_issues,
            "resume_reused_symbols": twelve_resume_reused_symbols,
            "refresh_symbols": sorted(refresh_symbols),
            "batches": twelve_batches,
            "symbol_failures": twelve_symbol_failures,
        },
        "tiingo": {
            "requested_symbols": tiingo_selected,
            "plan": tiingo_plan,
            "metadata": tiingo_metadata,
            "failures": tiingo_failures,
        },
        "eastmoney": {"queried": False, "role": settings["eastmoney_role"]},
        "candidate_files": candidate_files,
        "symbols": symbol_reports,
    }
    campaign_qualified, campaign_reasons = qualify_shadow_report(report, cfg)
    report["validation_campaign"] = {
        "qualified_session": campaign_qualified,
        "qualification_reasons": campaign_reasons,
    }
    report_content = json_bytes(report)
    markdown_content = shadow_markdown(report).encode("utf-8")
    atomic_write(processed_run / "report.json", report_content)
    atomic_write(processed_run / "report.md", markdown_content)
    latest = {
        "schema_version": 1,
        "run_id": run_id,
        "generated_at_utc": now.isoformat().replace("+00:00", "Z"),
        "desired_end_session": desired_end.isoformat(),
        "scope": report["scope"],
        "report_json": str((processed_run / "report.json").relative_to(WORKSPACE_ROOT)),
        "report_markdown": str((processed_run / "report.md").relative_to(WORKSPACE_ROOT)),
        "report_sha256": sha256_bytes(report_content),
        "summary": summary,
        "validation_campaign": report["validation_campaign"],
    }
    atomic_write(processed_root / "latest.json", json_bytes(latest))
    cross_check_issues = any(
        item.get("cross_check_status") in {"review", "failed"} for item in symbol_reports
    )
    exit_code = (
        0
        if summary["ready"] == summary["symbols"]
        and summary["fetch_failed"] == 0
        and not tiingo_failures
        and not cross_check_issues
        else 2
    )
    return report, exit_code


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    if not slug or len(slug) > 48:
        raise UpdateError("Provider/label must produce a 1-48 character lowercase slug")
    return slug


def safe_purchase_source(raw_path: str, destination_root: Path) -> Path:
    source = Path(raw_path).expanduser().resolve()
    forbidden = {Path("/").resolve(), Path.home().resolve(), WORKSPACE_ROOT.resolve(), (WORKSPACE_ROOT / "data").resolve()}
    if source in forbidden:
        raise UpdateError(f"Refusing an overly broad import source: {source}")
    if not source.exists():
        raise UpdateError(f"Purchase source does not exist: {source}")
    if source.is_symlink():
        raise UpdateError("Purchase source may not be a symlink")
    try:
        source.relative_to(WORKSPACE_ROOT.resolve())
    except ValueError:
        pass
    else:
        raise UpdateError("Purchased import source must be outside the managed Quant workspace")
    try:
        source.relative_to(destination_root.resolve())
    except ValueError:
        pass
    else:
        raise UpdateError("Purchase source is already inside the managed destination")
    return source


def purchase_files(source: Path) -> tuple[Path, list[Path]]:
    if source.is_file():
        return source.parent, [source]
    files: list[Path] = []
    for path in sorted(source.rglob("*")):
        if path.is_symlink():
            raise UpdateError(f"Symlinks are not allowed in purchased imports: {path}")
        if path.is_file() and path.name != ".DS_Store":
            files.append(path)
    if not files:
        raise UpdateError("Purchase source contains no files")
    return source, files


def inspect_file(path: Path, *, deep: bool) -> dict[str, Any]:
    suffix = path.suffix.lower()
    inspection: dict[str, Any] = {"suffix": suffix or "none", "status": "opaque"}
    try:
        if suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                bad = archive.testzip() if deep else None
                inspection = {
                    "suffix": suffix,
                    "status": "pass" if bad is None else "fail",
                    "archive_members": len(names),
                    "bad_member": bad,
                    "integrity_tested": deep,
                }
        elif suffix == ".csv":
            raw = path.read_bytes()[:65536]
            text = None
            encoding = None
            for candidate in ("utf-8-sig", "gb18030", "latin-1"):
                try:
                    text = raw.decode(candidate)
                    encoding = candidate
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise UpdateError("CSV header encoding could not be decoded")
            header = next(csv.reader(io.StringIO(text)), [])
            inspection = {"suffix": suffix, "status": "pass" if header else "fail", "encoding": encoding, "columns": header}
        elif suffix == ".json":
            with path.open(encoding="utf-8") as file:
                value = json.load(file)
            inspection = {"suffix": suffix, "status": "pass", "json_type": type(value).__name__}
        elif suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            inspection = {"suffix": suffix, "status": "pass", "sheets": workbook.sheetnames}
            workbook.close()
        elif suffix == ".xls":
            inspection = {"suffix": suffix, "status": "opaque", "note": "legacy XLS is inventoried but not parsed by intake"}
    except Exception as exc:
        inspection = {"suffix": suffix or "none", "status": "fail", "error": str(exc)}
    return inspection


def build_purchase_manifest(source: Path, *, deep: bool) -> dict[str, Any]:
    root, files = purchase_files(source)
    records = []
    digest = hashlib.sha256()
    for path in files:
        relative = path.name if source.is_file() else str(path.relative_to(root))
        file_hash = sha256_file(path)
        size = path.stat().st_size
        inspection = inspect_file(path, deep=deep)
        record = {
            "relative_path": relative,
            "bytes": size,
            "sha256": file_hash,
            "inspection": inspection,
        }
        records.append(record)
        digest.update(relative.encode("utf-8"))
        digest.update(b"\0")
        digest.update(str(size).encode("ascii"))
        digest.update(b"\0")
        digest.update(file_hash.encode("ascii"))
        digest.update(b"\n")
    failures = [item["relative_path"] for item in records if item["inspection"]["status"] == "fail"]
    return {
        "source_name": source.name,
        "root_hash": digest.hexdigest(),
        "file_count": len(records),
        "total_bytes": sum(item["bytes"] for item in records),
        "extension_counts": dict(Counter(Path(item["relative_path"]).suffix.lower() or "none" for item in records)),
        "inspection_failures": failures,
        "deep_inspection": deep,
        "files": records,
    }


def copy_purchase(
    source: Path,
    manifest: dict[str, Any],
    provider: str,
    acquired_date: str,
    label: str | None,
    cfg: dict[str, Any],
) -> dict[str, Any]:
    if manifest["inspection_failures"]:
        raise UpdateError(f"Cannot apply an import with inspection failures: {manifest['inspection_failures']}")
    provider_slug = slugify(provider)
    label_slug = slugify(label) if label else "batch"
    parsed_date = date.fromisoformat(acquired_date)
    destination_root = WORKSPACE_ROOT / cfg["purchased_import"]["destination_root"]
    batch_id = f"{parsed_date.isoformat()}_{label_slug}_{manifest['root_hash'][:12]}"
    target = destination_root / provider_slug / parsed_date.isoformat() / batch_id
    registry_path = WORKSPACE_ROOT / cfg["purchased_import"]["registry"]
    registry = load_json(registry_path)
    duplicate = next((item for item in registry.get("batches", []) if item["root_hash"] == manifest["root_hash"]), None)
    if duplicate:
        return {"status": "already_imported", "batch": duplicate}
    if target.exists():
        raise UpdateError(f"Unregistered destination already exists: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{batch_id}.", dir=str(target.parent)))
    try:
        source_root, source_files = purchase_files(source)
        by_relative = {
            (path.name if source.is_file() else str(path.relative_to(source_root))): path for path in source_files
        }
        for item in manifest["files"]:
            destination = staging / item["relative_path"]
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(by_relative[item["relative_path"]], destination, follow_symlinks=False)
            if sha256_file(destination) != item["sha256"]:
                raise UpdateError(f"Copied file hash mismatch: {item['relative_path']}")
            destination.chmod(stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
        os.replace(staging, target)
    except BaseException:
        if staging.exists():
            shutil.rmtree(staging)
        raise
    batch = {
        "batch_id": batch_id,
        "provider": provider,
        "acquired_date": parsed_date.isoformat(),
        "imported_at_utc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "source_name": manifest["source_name"],
        "destination": str(target.relative_to(WORKSPACE_ROOT)),
        "root_hash": manifest["root_hash"],
        "file_count": manifest["file_count"],
        "total_bytes": manifest["total_bytes"],
        "status": cfg["purchased_import"]["default_status"],
        "auto_merged": False,
        "files": manifest["files"],
    }
    registry.setdefault("batches", []).append(batch)
    atomic_write(registry_path, json_bytes(registry))
    return {"status": "imported_pending_review", "batch": batch}


def run_import(args: argparse.Namespace) -> tuple[dict[str, Any], int]:
    cfg = config()
    destination_root = WORKSPACE_ROOT / cfg["purchased_import"]["destination_root"]
    source = safe_purchase_source(args.source, destination_root)
    manifest = build_purchase_manifest(source, deep=args.deep or args.apply)
    result: dict[str, Any] = {
        "schema_version": 1,
        "mode": "apply" if args.apply else "preview",
        "source": manifest,
        "production_data_modified": False,
        "approved_data_modified": False,
    }
    if args.apply:
        if not args.provider or not args.acquired_date:
            raise UpdateError("--apply requires --provider and --acquired-date")
        result["apply"] = copy_purchase(source, manifest, args.provider, args.acquired_date, args.label, cfg)
    return result, 1 if manifest["inspection_failures"] else 0


def run_point_in_time_audit(args: argparse.Namespace) -> tuple[dict[str, Any], int]:
    report = audit_and_build(
        WORKSPACE_ROOT / args.registry,
        WORKSPACE_ROOT,
        write_candidate=bool(args.build_candidate),
    )
    return report, 1 if report.get("status") == "blocked_candidate_build" else 0


def print_import_human(payload: dict[str, Any]) -> None:
    source = payload["source"]
    print(f"Purchased data intake: {payload['mode']}")
    print(f"- source: {source['source_name']}")
    print(f"- files: {source['file_count']} / bytes: {source['total_bytes']}")
    print(f"- root SHA256: {source['root_hash']}")
    print(f"- inspection failures: {len(source['inspection_failures'])}")
    if payload["mode"] == "preview":
        print("- no files copied; rerun with explicit --apply, --provider, and --acquired-date after reviewing this plan")
    else:
        applied = payload["apply"]
        print(f"- result: {applied['status']}")
        print(f"- destination: {applied['batch']['destination']}")
        print("- status: pending_review; not merged into approved data")


def print_point_in_time_human(payload: dict[str, Any]) -> None:
    summary = payload["summary"]
    point_in_time = payload["point_in_time"]
    comparison = payload["package_comparison"]
    print(f"{payload.get('index_display_name', 'Index')} point-in-time source audit")
    print(f"- status: {payload['status']}; approved=false; production writes=0")
    print(
        f"- instruments/rows/intervals: {summary['instrument_count']}/"
        f"{summary['price_row_count']}/{summary['membership_interval_count']}"
    )
    print(
        f"- membership: {point_in_time['first_membership_date']} to {point_in_time['last_membership_date']}; "
        f"daily count {point_in_time['minimum_member_count']} to {point_in_time['maximum_member_count']}"
    )
    print(
        f"- package identity fields: key mismatches={comparison['package_row_key_mismatch']}; "
        f"InIndex mismatches={comparison['package_in_index_mismatch']}; "
        f"Unadjusted Close mismatches={comparison['package_unadjusted_close_mismatch']}"
    )
    print(
        f"- official cross-checks: {summary['official_cross_checks_passed']}/"
        f"{summary['official_cross_checks_total']} passed"
    )
    print(f"- blocking issues: {payload['blocking_issue_counts'] or 'none'}")
    print(f"- candidate outputs written: {str(payload.get('candidate_outputs_written', False)).lower()}")
    if payload.get("output_directory"):
        print(f"- output directory: {payload['output_directory']}")


def add_json_flag(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    status = subparsers.add_parser("status", help="Show database freshness and operational state")
    add_json_flag(status)
    check = subparsers.add_parser("check", help="Check database and evidence integrity")
    check.add_argument("--deep", action="store_true", help="Verify frozen/imported hashes and run raw inventory, audit, and all tests")
    add_json_flag(check)
    shadow = subparsers.add_parser("shadow-update", help="Fetch isolated current-member candidates without production writes")
    shadow.add_argument("--symbols", help="Comma-separated subset of current SPY equities")
    shadow.add_argument("--limit", type=int, help="Limit the deterministic current-member list")
    shadow.add_argument("--end", help="Target date; resolved to the latest XNYS session on or before it")
    shadow.add_argument("--run-id", help="Explicit immutable run ID")
    shadow.add_argument(
        "--replay-twelve-run",
        help="Reanalyze immutable Twelve Data archives from an earlier run without spending credits",
    )
    shadow.add_argument(
        "--resume-twelve-run",
        help="Reuse a partial immutable Twelve Data archive and fetch only missing or stale symbols",
    )
    shadow.add_argument(
        "--refresh-symbols",
        help="Comma-separated selected symbols to refetch while resuming instead of reusing their archived series",
    )
    shadow.add_argument(
        "--replay-tiingo-runs",
        help="Comma-separated immutable run IDs; reuse archived Tiingo payloads and fetch only missing symbols",
    )
    shadow.add_argument("--tiingo", choices=("rotate", "all", "none"), default="rotate")
    add_json_flag(shadow)
    intake = subparsers.add_parser("import-purchased", help="Preview or explicitly apply an immutable purchased-data intake")
    intake.add_argument("--source", required=True, help="Exact source file or directory")
    intake.add_argument("--provider", help="Provider name; required with --apply")
    intake.add_argument("--acquired-date", help="YYYY-MM-DD; required with --apply")
    intake.add_argument("--label", help="Short batch label")
    intake.add_argument("--deep", action="store_true", help="Run ZIP integrity tests during preview")
    intake.add_argument("--apply", action="store_true", help="Copy into immutable raw intake after validation")
    add_json_flag(intake)
    nasdaq100 = subparsers.add_parser(
        "audit-nasdaq100",
        help="Audit purchased Nasdaq-100 point-in-time evidence without using paid APIs",
    )
    nasdaq100.add_argument(
        "--registry",
        default="data/nasdaq100_history_registry.json",
        help="Workspace-relative candidate registry path",
    )
    nasdaq100.add_argument(
        "--build-candidate",
        action="store_true",
        help="Write compact pending_review products after all blocking checks pass",
    )
    add_json_flag(nasdaq100)
    sp500 = subparsers.add_parser(
        "audit-sp500",
        help="Reaudit purchased S&P 500 point-in-time evidence without using paid APIs",
    )
    sp500.add_argument(
        "--registry",
        default="data/sp500_candidate_registry.json",
        help="Workspace-relative candidate registry path",
    )
    sp500.add_argument(
        "--build-candidate",
        action="store_true",
        help="Write compact pending_review products after all blocking checks pass",
    )
    add_json_flag(sp500)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args.command == "status":
            payload = build_status()
            if args.json:
                print(json.dumps(payload, ensure_ascii=False, indent=2))
            else:
                print_status_human(payload)
            return 0
        if args.command == "check":
            payload = quick_check(verify_hashes=args.deep)
            if args.deep and payload["status"] != "fail":
                payload["process_checks"] = run_deep_process_checks()
                if any(item["status"] == "fail" for item in payload["process_checks"]):
                    payload["status"] = "fail"
            if args.json:
                print(json.dumps(payload, ensure_ascii=False, indent=2))
            else:
                print_check_human(payload)
            return 1 if payload["status"] == "fail" else 0
        if args.command == "shadow-update":
            payload, exit_code = run_shadow_update(args)
            if args.json:
                print(json.dumps(payload, ensure_ascii=False, indent=2))
            else:
                print(json.dumps({"run_id": payload["run_id"], "summary": payload["summary"]}, ensure_ascii=False))
            return exit_code
        if args.command == "import-purchased":
            payload, exit_code = run_import(args)
            if args.json:
                print(json.dumps(payload, ensure_ascii=False, indent=2))
            else:
                print_import_human(payload)
            return exit_code
        if args.command in {"audit-nasdaq100", "audit-sp500"}:
            payload, exit_code = run_point_in_time_audit(args)
            if args.json:
                print(json.dumps(payload, ensure_ascii=False, indent=2))
            else:
                print_point_in_time_human(payload)
            return exit_code
    except (UpdateError, PointInTimeDataError, ProviderError, OSError, ValueError, zipfile.BadZipFile) as exc:
        print(f"data-update error: {exc}", file=sys.stderr)
        return 2
    raise AssertionError(args.command)


if __name__ == "__main__":
    raise SystemExit(main())
