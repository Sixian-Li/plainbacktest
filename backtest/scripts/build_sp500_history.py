#!/usr/bin/env python3
"""Build the point-in-time S&P 500 history baseline without changing raw data."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import platform
import shutil
import tempfile
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import exchange_calendars as xcals
import openpyxl
from openpyxl import load_workbook


WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_REGISTRY = WORKSPACE_ROOT / "data/sp500_history_registry.json"
CANONICAL_COLUMNS = ("date", "symbol", "open", "high", "low", "close", "volume")
SOURCE_COLUMNS = (
    "Date",
    "CompanyName",
    "Symbol",
    "Open",
    "High",
    "Low",
    "Close",
    "Volume",
    "Unadjusted Close",
    "InIndex",
)
PRICE_COLUMNS = ("Open", "High", "Low", "Close")
MEMBERSHIP_COLUMNS = (
    "instrument_id",
    "symbol",
    "run_number",
    "observed_start",
    "observed_end",
    "previous_observed_date",
    "next_observed_date",
    "start_boundary",
    "end_boundary",
    "start_gap_sessions",
    "end_gap_sessions",
    "entry_observed",
    "exit_observed",
    "source_snapshot_date",
)
SECURITY_MASTER_COLUMNS = (
    "instrument_id",
    "symbol",
    "company_name",
    "source_member",
    "price_start",
    "price_end",
    "price_rows",
    "membership_runs",
    "first_row_in_index",
    "last_row_in_index",
    "active_on_vendor_snapshot",
    "current_official_member",
    "current_official_identifier",
    "current_official_sedol",
    "current_link_method",
    "quality_status",
    "member_price_coverage_status",
)
CURRENT_COLUMNS = (
    "as_of_date",
    "symbol",
    "company_name",
    "identifier",
    "sedol",
    "weight_pct",
    "shares_held",
    "local_currency",
    "vendor_instrument_id",
    "vendor_price_available",
    "link_method",
)
RECONCILIATION_COLUMNS = (
    "symbol",
    "vendor_snapshot_date",
    "vendor_active",
    "official_snapshot_date",
    "official_active",
    "classification",
)
EXCLUSION_COLUMNS = (
    "as_of_date",
    "ticker",
    "name",
    "identifier",
    "weight_pct",
    "classification",
    "reason",
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=str(path.parent))
    try:
        with os.fdopen(descriptor, "wb") as file:
            file.write(data)
        os.replace(temp_name, path)
    except BaseException:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def format_number(value: float) -> str:
    return format(value, ".15g")


def csv_bytes(fieldnames: Iterable[str], rows: Iterable[dict[str, Any]]) -> bytes:
    text = io.StringIO(newline="")
    writer = csv.DictWriter(text, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    for row in rows:
        values = {}
        for field in fieldnames:
            value = row.get(field, "")
            values[field] = str(value).lower() if isinstance(value, bool) else value
        writer.writerow(values)
    return text.getvalue().encode("utf-8")


def issue_record(
    issue_counts: Counter[str],
    issue_examples: dict[str, list[dict[str, Any]]],
    maximum_examples: int,
    name: str,
    payload: dict[str, Any],
) -> None:
    issue_counts[name] += 1
    examples = issue_examples.setdefault(name, [])
    if len(examples) < maximum_examples:
        examples.append(payload)


def parse_finite(raw: Any, field: str) -> float:
    text = "" if raw is None else str(raw).strip()
    if not text:
        raise ValueError(f"missing {field}")
    value = float(text)
    if not math.isfinite(value):
        raise ValueError(f"non-finite {field}={text!r}")
    return value


def parse_holdings(registry: dict[str, Any], workspace: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    source = registry["current_holdings_source"]
    source_path = workspace / source["path"]
    actual_hash = sha256_file(source_path)
    if actual_hash != source["sha256"]:
        raise ValueError(f"State Street snapshot SHA256 mismatch: {actual_hash}")

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook["holdings"]
    if sheet["A1"].value != "Fund Name:" or sheet["B2"].value != "SPY":
        raise ValueError("Unexpected State Street workbook identity")
    as_of_text = str(sheet["B3"].value or "")
    parsed_as_of = datetime.strptime(source["as_of_date"], "%Y-%m-%d")
    english_months = (
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    )
    expected_text = f"As of {parsed_as_of.day:02d}-{english_months[parsed_as_of.month - 1]}-{parsed_as_of.year}"
    if as_of_text != expected_text:
        raise ValueError(f"Unexpected holdings date: {as_of_text!r}; expected {expected_text!r}")
    headers = [cell.value for cell in sheet[5]][:8]
    expected_headers = [
        "Name",
        "Ticker",
        "Identifier",
        "SEDOL",
        "Weight",
        "Sector",
        "Shares Held",
        "Local Currency",
    ]
    if headers != expected_headers:
        raise ValueError(f"Unexpected holdings headers: {headers!r}")

    exclusion_rules = {item["ticker"]: item for item in source["excluded_holdings"]}
    equities: list[dict[str, Any]] = []
    exclusions: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=6, values_only=True):
        name, ticker, identifier, sedol, weight, _sector, shares, currency = values[:8]
        if ticker is None:
            break
        ticker = str(ticker).strip()
        base = {
            "as_of_date": source["as_of_date"],
            "symbol": ticker,
            "company_name": str(name or "").strip(),
            "identifier": str(identifier or "").strip(),
            "sedol": str(sedol or "").strip(),
            "weight_pct": format_number(float(weight)),
            "shares_held": format_number(float(shares)),
            "local_currency": str(currency or "").strip(),
        }
        if ticker in exclusion_rules:
            rule = exclusion_rules[ticker]
            exclusions.append(
                {
                    "as_of_date": base["as_of_date"],
                    "ticker": ticker,
                    "name": base["company_name"],
                    "identifier": base["identifier"],
                    "weight_pct": base["weight_pct"],
                    "classification": rule["classification"],
                    "reason": rule["reason"],
                }
            )
        else:
            equities.append(base)
    workbook.close()
    if len(equities) + len(exclusions) != registry["expected"]["official_snapshot_raw_holdings"]:
        raise ValueError("Unexpected number of raw State Street holdings")
    return equities, exclusions


def membership_intervals(
    instrument_id: str,
    symbol: str,
    rows: list[dict[str, Any]],
    snapshot_date: str,
    session_positions: dict[str, int],
) -> list[dict[str, Any]]:
    intervals: list[dict[str, Any]] = []
    run_start: int | None = None
    for index, row in enumerate(rows):
        active = row["in_index"] == 1
        if active and run_start is None:
            run_start = index
        at_end = index == len(rows) - 1
        if run_start is not None and ((not active) or at_end):
            run_end = index if active and at_end else index - 1
            left_censored = run_start == 0
            right_censored = run_end == len(rows) - 1
            previous_observed_date = "" if left_censored else rows[run_start - 1]["date"]
            next_observed_date = "" if right_censored else rows[run_end + 1]["date"]
            start_gap_sessions = (
                ""
                if left_censored
                else max(
                    0,
                    session_positions[rows[run_start]["date"]]
                    - session_positions[previous_observed_date]
                    - 1,
                )
            )
            end_gap_sessions = (
                ""
                if right_censored
                else max(
                    0,
                    session_positions[next_observed_date]
                    - session_positions[rows[run_end]["date"]]
                    - 1,
                )
            )
            if left_censored:
                start_boundary = "left_censored"
                entry_observed = False
            elif start_gap_sessions == 0:
                start_boundary = "adjacent_observed_0_to_1"
                entry_observed = True
            else:
                start_boundary = "interval_censored_0_to_1"
                entry_observed = False
            if right_censored:
                end_boundary = "right_censored"
                exit_observed = False
            elif end_gap_sessions == 0:
                end_boundary = "adjacent_observed_1_to_0"
                exit_observed = True
            else:
                end_boundary = "interval_censored_1_to_0"
                exit_observed = False
            intervals.append(
                {
                    "instrument_id": instrument_id,
                    "symbol": symbol,
                    "run_number": len(intervals) + 1,
                    "observed_start": rows[run_start]["date"],
                    "observed_end": rows[run_end]["date"],
                    "previous_observed_date": previous_observed_date,
                    "next_observed_date": next_observed_date,
                    "start_boundary": start_boundary,
                    "end_boundary": end_boundary,
                    "start_gap_sessions": start_gap_sessions,
                    "end_gap_sessions": end_gap_sessions,
                    "entry_observed": str(entry_observed).lower(),
                    "exit_observed": str(exit_observed).lower(),
                    "source_snapshot_date": snapshot_date,
                }
            )
            run_start = None
    return intervals


def parse_member(
    member: str,
    source_bytes: bytes,
    sessions: set[str],
    session_positions: dict[str, int],
    registry: dict[str, Any],
) -> tuple[dict[str, Any], list[dict[str, Any]], bytes, dict[str, Any]]:
    policy = registry["quality_policy"]
    maximum_examples = int(policy["max_issue_examples"])
    issue_counts: Counter[str] = Counter()
    issue_examples: dict[str, list[dict[str, Any]]] = {}
    instrument_id = Path(member).stem
    if Path(member).name != member or not member.endswith(".csv") or not instrument_id:
        issue_record(issue_counts, issue_examples, maximum_examples, "bad_archive_member", {"member": member})

    reader = csv.DictReader(io.StringIO(source_bytes.decode("utf-8-sig"), newline=""))
    if tuple(reader.fieldnames or ()) != SOURCE_COLUMNS:
        issue_record(
            issue_counts,
            issue_examples,
            maximum_examples,
            "unexpected_schema",
            {"actual": reader.fieldnames, "expected": list(SOURCE_COLUMNS)},
        )

    parsed_rows: list[dict[str, Any]] = []
    company_names: set[str] = set()
    seen_dates: set[str] = set()
    previous_date: str | None = None
    previous_close: float | None = None
    previous_in_index: int | None = None
    previous_factor: float | None = None
    source_rows = 0
    for row_number, source_row in enumerate(reader, start=2):
        source_rows += 1
        missing = [column for column in SOURCE_COLUMNS if not str(source_row.get(column) or "").strip()]
        if missing:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "missing_required_value",
                {"row": row_number, "columns": missing},
            )
            continue
        date_text = str(source_row["Date"]).strip()
        try:
            date.fromisoformat(date_text)
        except ValueError:
            issue_record(
                issue_counts, issue_examples, maximum_examples, "bad_date", {"row": row_number, "value": date_text}
            )
            continue

        symbol = str(source_row["Symbol"]).strip().upper()
        company_name = str(source_row["CompanyName"]).strip()
        company_names.add(company_name)
        if symbol != instrument_id:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "wrong_symbol",
                {"row": row_number, "date": date_text, "actual": symbol, "expected": instrument_id},
            )
        if date_text in seen_dates:
            issue_record(
                issue_counts, issue_examples, maximum_examples, "duplicate_date", {"row": row_number, "date": date_text}
            )
        seen_dates.add(date_text)
        prior_date = previous_date
        if prior_date is not None and date_text <= prior_date:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "source_order_violation",
                {"row": row_number, "previous": prior_date, "date": date_text},
            )
        previous_date = date_text
        if date_text not in sessions:
            issue_record(
                issue_counts, issue_examples, maximum_examples, "non_session_date", {"row": row_number, "date": date_text}
            )

        try:
            numeric = {
                field: parse_finite(source_row[field], field)
                for field in (*PRICE_COLUMNS, "Volume", "Unadjusted Close")
            }
        except (ValueError, TypeError) as exc:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "bad_number",
                {"row": row_number, "date": date_text, "detail": str(exc)},
            )
            continue
        prices = [numeric[field] for field in PRICE_COLUMNS]
        if min(prices) <= 0:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "nonpositive_price",
                {"row": row_number, "date": date_text, "values": prices},
            )
        tolerance = float(policy["ohlc_relative_tolerance"]) * max(1.0, *(abs(value) for value in prices))
        if numeric["High"] + tolerance < max(prices) or numeric["Low"] - tolerance > min(prices):
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "invalid_ohlc",
                {
                    "row": row_number,
                    "date": date_text,
                    "open": numeric["Open"],
                    "high": numeric["High"],
                    "low": numeric["Low"],
                    "close": numeric["Close"],
                },
            )
        if numeric["Volume"] < 0:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "negative_volume",
                {"row": row_number, "date": date_text, "volume": numeric["Volume"]},
            )
        elif numeric["Volume"] == 0:
            issue_record(
                issue_counts, issue_examples, maximum_examples, "zero_volume", {"row": row_number, "date": date_text}
            )
        if numeric["Unadjusted Close"] <= 0:
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "nonpositive_unadjusted_close",
                {"row": row_number, "date": date_text, "value": numeric["Unadjusted Close"]},
            )
        try:
            in_index = int(str(source_row["InIndex"]).strip())
        except ValueError:
            in_index = -1
        if in_index not in (0, 1):
            issue_record(
                issue_counts,
                issue_examples,
                maximum_examples,
                "bad_in_index",
                {"row": row_number, "date": date_text, "value": source_row["InIndex"]},
            )

        if previous_close is not None and previous_close > 0:
            daily_return = numeric["Close"] / previous_close - 1.0
            if abs(daily_return) > float(policy["outlier_abs_close_return_pct"]) / 100.0:
                issue_record(
                    issue_counts,
                    issue_examples,
                    maximum_examples,
                    "outlier_close_return",
                    {"date": date_text, "return_pct": round(daily_return * 100.0, 9)},
                )
                if (
                    previous_in_index == 1
                    and in_index == 1
                    and prior_date in session_positions
                    and date_text in session_positions
                    and session_positions[date_text] - session_positions[prior_date] == 1
                ):
                    issue_record(
                        issue_counts,
                        issue_examples,
                        maximum_examples,
                        "adjacent_in_index_outlier_close_return",
                        {"date": date_text, "return_pct": round(daily_return * 100.0, 9)},
                    )
        previous_close = numeric["Close"]
        previous_in_index = in_index
        factor = numeric["Close"] / numeric["Unadjusted Close"]
        if previous_factor is not None and previous_factor > 0:
            factor_change = factor / previous_factor - 1.0
            if abs(factor_change) > float(policy["adjustment_factor_change_tolerance"]):
                issue_record(
                    issue_counts,
                    issue_examples,
                    maximum_examples,
                    "adjustment_factor_change",
                    {"date": date_text, "change_pct": round(factor_change * 100.0, 9)},
                )
        previous_factor = factor
        parsed_rows.append(
            {
                "date": date_text,
                "symbol": symbol,
                "open": numeric["Open"],
                "high": numeric["High"],
                "low": numeric["Low"],
                "close": numeric["Close"],
                "volume": numeric["Volume"],
                "in_index": in_index,
            }
        )

    if len(company_names) != 1:
        issue_record(
            issue_counts,
            issue_examples,
            maximum_examples,
            "multiple_company_names",
            {"count": len(company_names), "names": sorted(company_names)[:maximum_examples]},
        )
    intervals = membership_intervals(
        instrument_id,
        instrument_id,
        parsed_rows,
        registry["price_source"]["snapshot_date"],
        session_positions,
    )
    if not intervals:
        issue_record(issue_counts, issue_examples, maximum_examples, "never_in_index", {})

    if parsed_rows:
        expected_sessions = {day for day in sessions if parsed_rows[0]["date"] <= day <= parsed_rows[-1]["date"]}
        missing_sessions = sorted(expected_sessions - seen_dates)
        issue_counts["missing_session"] = len(missing_sessions)
        if missing_sessions:
            issue_examples["missing_session"] = [{"date": day} for day in missing_sessions[:maximum_examples]]
        missing_in_index_sessions: list[str] = []
        for interval in intervals:
            expected_interval = {
                day
                for day in sessions
                if interval["observed_start"] <= day <= interval["observed_end"]
            }
            missing_in_index_sessions.extend(sorted(expected_interval - seen_dates))
        issue_counts["missing_in_index_session"] = len(missing_in_index_sessions)
        if missing_in_index_sessions:
            issue_examples["missing_in_index_session"] = [
                {"date": day} for day in missing_in_index_sessions[:maximum_examples]
            ]

    blocking = {
        name: issue_counts[name]
        for name in policy["blocking_checks"]
        if issue_counts[name]
    }
    quality_status = "failed" if blocking else "passed"
    price_rows = [
        {
            "date": row["date"],
            "symbol": row["symbol"],
            "open": format_number(row["open"]),
            "high": format_number(row["high"]),
            "low": format_number(row["low"]),
            "close": format_number(row["close"]),
            "volume": format_number(row["volume"]),
        }
        for row in parsed_rows
    ]
    output_bytes = csv_bytes(CANONICAL_COLUMNS, price_rows)
    metadata = {
        "instrument_id": instrument_id,
        "symbol": instrument_id,
        "company_name": next(iter(company_names), ""),
        "source_member": member,
        "source_member_bytes": len(source_bytes),
        "source_member_sha256": sha256_bytes(source_bytes),
        "price_start": parsed_rows[0]["date"] if parsed_rows else None,
        "price_end": parsed_rows[-1]["date"] if parsed_rows else None,
        "source_rows": source_rows,
        "price_rows": len(parsed_rows),
        "canonical_sha256": sha256_bytes(output_bytes),
        "membership_runs": len(intervals),
        "first_row_in_index": bool(parsed_rows and parsed_rows[0]["in_index"] == 1),
        "last_row_in_index": bool(parsed_rows and parsed_rows[-1]["in_index"] == 1),
        "active_on_vendor_snapshot": bool(
            parsed_rows
            and parsed_rows[-1]["date"] == registry["price_source"]["snapshot_date"]
            and parsed_rows[-1]["in_index"] == 1
        ),
        "quality_status": quality_status,
        "member_price_coverage_status": (
            "gaps_reported" if issue_counts["missing_in_index_session"] else "complete"
        ),
        "blocking_counts": blocking,
        "issue_counts": {name: issue_counts[name] for name in sorted(issue_counts) if issue_counts[name]},
        "issue_examples": {name: issue_examples[name] for name in sorted(issue_examples)},
    }
    return metadata, intervals, output_bytes, {
        "instrument_id": instrument_id,
        "quality_status": quality_status,
        "blocking_counts": blocking,
        "issue_counts": metadata["issue_counts"],
        "issue_examples": metadata["issue_examples"],
    }


def render_quality_markdown(
    manifest: dict[str, Any],
    quality: dict[str, Any],
    registry: dict[str, Any],
) -> str:
    summary = quality["summary"]
    reconciliation = manifest["snapshot_reconciliation"]
    lines = [
        "# S&P 500 历史基线质量报告",
        "",
        f"- 构建 ID：`{manifest['build_id']}`",
        f"- 历史价格截止：`{registry['price_source']['snapshot_date']}`",
        f"- 官方 SPY 持仓快照：`{registry['current_holdings_source']['as_of_date']}`",
        f"- 历史证券：{summary['instrument_count']:,} 个，价格 {summary['price_row_count']:,} 行",
        f"- 成员区间：{summary['membership_interval_count']:,} 段",
        f"- 价格质量：{summary['passed_instruments']:,} 通过 / {summary['failed_instruments']:,} 失败",
        f"- 成员期价格有缺口：{summary['instruments_with_member_price_gaps']:,} 个证券（不填值）",
        "",
        "## 快照对账",
        "",
        f"- 供应商 `{reconciliation['vendor_snapshot_date']}` 成分股：{reconciliation['vendor_active_count']} 只",
        f"- State Street `{reconciliation['official_snapshot_date']}` 普通股：{reconciliation['official_active_count']} 只",
        f"- 两快照共同：{reconciliation['matched_count']} 只",
        f"- 仅最新官方快照：{', '.join(reconciliation['official_only']) or '无'}",
        f"- 仅历史包截止快照：{', '.join(reconciliation['vendor_only']) or '无'}",
        "",
        "快照日期不同，差异只表示期间成分变化；不自动推断删除原因。",
        "",
        "## 成员边界",
        "",
        f"- 左侧截断区间：{summary['left_censored_intervals']:,} 段",
        f"- 右侧截断区间：{summary['right_censored_intervals']:,} 段",
        f"- 多段成员证券：{summary['multi_run_instruments']:,} 个",
        f"- 跨越缺失 session 的入选/退出边界：{summary['entry_boundaries_with_missing_sessions']:,} / {summary['exit_boundaries_with_missing_sessions']:,}",
        "",
        "`右侧截断` 不等于当前仍在指数内；当前范围只由带日期的 State Street 快照决定。边界前后若有缺失 session，只能确定变化发生在两个可观察日之间。",
        "",
        "## 价格检查",
        "",
    ]
    if summary["blocking_issue_counts"]:
        lines.append("### 阻断问题")
        lines.append("")
        for name, count in summary["blocking_issue_counts"].items():
            lines.append(f"- `{name}`: {count:,}")
    else:
        lines.append("- 所有历史证券均通过结构、数值、OHLC、美股交易日及成员标记阻断检查。")
    lines.extend(["", "非阻断观察：", ""])
    warning_labels = {
        "missing_session": "单证券缺少的交易日",
        "missing_in_index_session": "可观察成员期内缺少的价格交易日",
        "zero_volume": "零成交量行",
        "outlier_close_return": "绝对收盘变动超过 30% 的行",
        "adjacent_in_index_outlier_close_return": "成员期内相邻 session 绝对收盘变动超过 30% 的行",
        "adjustment_factor_change": "相邻日复权比率变化超过容差的行",
    }
    for name, count in summary["warning_issue_counts"].items():
        instruments = summary["instruments_with_warning"][name]
        lines.append(
            f"- `{name}`（{warning_labels[name]}）：{count:,} 条，涉及 {instruments:,} 个证券"
        )
    lines.extend(
        [
            "",
            "这些是供应商原值中的审计信号，不会被填值、删除或平滑；其中复权比率变化包含正常的分红/拆股事件，不单独视为错误。",
        ]
    )
    lines.extend(["", "## 使用规则", ""])
    lines.extend(f"- {item}" for item in registry["known_limitations"])
    lines.extend(
        [
            "- 历史成员信号默认滞后一个交易日才可使用，除非另有公告时间证据。",
            "- 成员期内的缺失价格不得前向填充；回测必须将该证券当日视为不可交易并报告。",
            "- 更新器不得把新快照直接覆盖旧快照，也不得仅追加复权价格而不比较重叠窗口。",
            "",
        ]
    )
    return "\n".join(lines)


def build(registry_path: Path, workspace: Path, sample: set[str] | None = None) -> dict[str, Any]:
    registry_bytes = registry_path.read_bytes()
    registry = json.loads(registry_bytes.decode("utf-8"))
    if registry.get("schema_version") != 1:
        raise ValueError("Unsupported registry schema_version")

    for frozen in registry["frozen_existing_outputs"]:
        actual = sha256_file(workspace / frozen["path"])
        if actual != frozen["sha256"]:
            raise ValueError(f"Frozen existing output changed: {frozen['path']}")
    source_path = workspace / registry["price_source"]["path"]
    actual_archive_hash = sha256_file(source_path)
    if actual_archive_hash != registry["price_source"]["sha256"]:
        raise ValueError(f"Price archive SHA256 mismatch: {actual_archive_hash}")

    official_equities, exclusions = parse_holdings(registry, workspace)
    official_by_symbol = {row["symbol"]: row for row in official_equities}
    if len(official_by_symbol) != len(official_equities):
        raise ValueError("Duplicate equity ticker in official snapshot")

    calendar = xcals.get_calendar("XNYS", start="1990-01-01", end="2026-12-31")
    sessions = set(calendar.sessions.strftime("%Y-%m-%d"))
    session_positions = {day: index for index, day in enumerate(sorted(sessions))}
    output_config = registry["outputs"]
    price_root = workspace / output_config["price_directory"]
    sample_mode = sample is not None
    if sample_mode:
        temporary_root = Path(tempfile.mkdtemp(prefix="sp500-history-sample-"))
        price_target_root = temporary_root / "prices"
    else:
        price_target_root = price_root.parent / f".{price_root.name}.build"
        if price_target_root.exists():
            shutil.rmtree(price_target_root)
    price_target_root.mkdir(parents=True, exist_ok=True)

    instruments: list[dict[str, Any]] = []
    intervals: list[dict[str, Any]] = []
    reports: list[dict[str, Any]] = []
    with zipfile.ZipFile(source_path) as archive:
        all_members = sorted(archive.namelist())
        members = [member for member in all_members if sample is None or Path(member).stem in sample]
        missing_sample = sorted((sample or set()) - {Path(member).stem for member in members})
        if missing_sample:
            raise ValueError(f"Sample instruments absent from archive: {missing_sample}")
        for index, member in enumerate(members, start=1):
            source_bytes = archive.read(member)
            metadata, member_intervals, output_bytes, report = parse_member(
                member, source_bytes, sessions, session_positions, registry
            )
            atomic_write(price_target_root / member, output_bytes)
            instruments.append(metadata)
            intervals.extend(member_intervals)
            reports.append(report)
            if not sample_mode and (index % 100 == 0 or index == len(members)):
                print(f"Parsed {index}/{len(members)} archive members", flush=True)

    if sample_mode:
        result = {
            "sample": sorted(sample or set()),
            "instruments": instruments,
            "membership_intervals": intervals,
            "reports": reports,
            "temporary_output": str(price_target_root),
        }
        shutil.rmtree(temporary_root)
        return result

    vendor_active = {item["symbol"] for item in instruments if item["active_on_vendor_snapshot"]}
    official_active = set(official_by_symbol)
    matched = vendor_active & official_active
    official_only = official_active - vendor_active
    vendor_only = vendor_active - official_active

    current_rows: list[dict[str, Any]] = []
    for symbol in sorted(official_active):
        official = official_by_symbol[symbol]
        active_vendor_metadata = next(
            (item for item in instruments if item["instrument_id"] == symbol and item["active_on_vendor_snapshot"]),
            None,
        )
        current_rows.append(
            {
                **official,
                "vendor_instrument_id": active_vendor_metadata["instrument_id"] if active_vendor_metadata else "",
                "vendor_price_available": str(active_vendor_metadata is not None).lower(),
                "link_method": "exact_ticker_adjacent_snapshot" if active_vendor_metadata else "unmatched",
            }
        )

    for item in instruments:
        official = official_by_symbol.get(item["symbol"]) if item["active_on_vendor_snapshot"] else None
        item["current_official_member"] = official is not None
        item["current_official_identifier"] = official["identifier"] if official else ""
        item["current_official_sedol"] = official["sedol"] if official else ""
        item["current_link_method"] = "exact_ticker_adjacent_snapshot" if official else ""

    reconciliation_rows = []
    for symbol in sorted(vendor_active | official_active):
        in_vendor = symbol in vendor_active
        in_official = symbol in official_active
        reconciliation_rows.append(
            {
                "symbol": symbol,
                "vendor_snapshot_date": registry["price_source"]["snapshot_date"],
                "vendor_active": str(in_vendor).lower(),
                "official_snapshot_date": registry["current_holdings_source"]["as_of_date"],
                "official_active": str(in_official).lower(),
                "classification": (
                    "matched" if in_vendor and in_official else "official_only" if in_official else "vendor_only"
                ),
            }
        )

    expected = registry["expected"]
    actual_expectations = {
        "archive_members": len(instruments),
        "archive_price_rows": sum(item["price_rows"] for item in instruments),
        "ever_in_index_members": sum(item["membership_runs"] > 0 for item in instruments),
        "vendor_snapshot_active_equities": len(vendor_active),
        "official_snapshot_raw_holdings": len(official_equities) + len(exclusions),
        "official_snapshot_equities": len(official_active),
        "matched_active_equities": len(matched),
        "official_snapshot_only": sorted(official_only),
        "vendor_snapshot_only": sorted(vendor_only),
        "coverage_not_before": min(item["price_start"] for item in instruments),
    }
    expectation_mismatches = {
        key: {"expected": expected[key], "actual": actual}
        for key, actual in actual_expectations.items()
        if actual != expected[key]
    }
    if expectation_mismatches:
        raise ValueError(f"Baseline expectation mismatch: {json.dumps(expectation_mismatches, ensure_ascii=False)}")

    failed_instruments = [item for item in instruments if item["quality_status"] == "failed"]
    blocking_issue_counts: Counter[str] = Counter()
    warning_issue_counts: Counter[str] = Counter()
    instruments_with_warning: Counter[str] = Counter()
    for report in reports:
        blocking_issue_counts.update(report["blocking_counts"])
        for name in registry["quality_policy"]["warning_checks"]:
            count = report["issue_counts"].get(name, 0)
            warning_issue_counts[name] += count
            if count:
                instruments_with_warning[name] += 1
    summary = {
        "instrument_count": len(instruments),
        "price_row_count": actual_expectations["archive_price_rows"],
        "membership_interval_count": len(intervals),
        "passed_instruments": len(instruments) - len(failed_instruments),
        "failed_instruments": len(failed_instruments),
        "left_censored_intervals": sum(item["start_boundary"] == "left_censored" for item in intervals),
        "right_censored_intervals": sum(item["end_boundary"] == "right_censored" for item in intervals),
        "multi_run_instruments": sum(item["membership_runs"] > 1 for item in instruments),
        "entry_boundaries_with_missing_sessions": sum(
            isinstance(item["start_gap_sessions"], int) and item["start_gap_sessions"] > 0
            for item in intervals
        ),
        "exit_boundaries_with_missing_sessions": sum(
            isinstance(item["end_gap_sessions"], int) and item["end_gap_sessions"] > 0
            for item in intervals
        ),
        "instruments_with_member_price_gaps": sum(
            item["member_price_coverage_status"] == "gaps_reported" for item in instruments
        ),
        "blocking_issue_counts": dict(sorted(blocking_issue_counts.items())),
        "warning_issue_counts": {
            name: warning_issue_counts[name]
            for name in registry["quality_policy"]["warning_checks"]
        },
        "instruments_with_warning": {
            name: instruments_with_warning[name]
            for name in registry["quality_policy"]["warning_checks"]
        },
    }
    reconciliation_summary = {
        "vendor_snapshot_date": registry["price_source"]["snapshot_date"],
        "vendor_active_count": len(vendor_active),
        "official_snapshot_date": registry["current_holdings_source"]["as_of_date"],
        "official_active_count": len(official_active),
        "matched_count": len(matched),
        "official_only": sorted(official_only),
        "vendor_only": sorted(vendor_only),
    }

    security_master_bytes = csv_bytes(SECURITY_MASTER_COLUMNS, sorted(instruments, key=lambda row: row["instrument_id"]))
    interval_bytes = csv_bytes(
        MEMBERSHIP_COLUMNS,
        sorted(intervals, key=lambda row: (row["instrument_id"], int(row["run_number"]))),
    )
    current_bytes = csv_bytes(CURRENT_COLUMNS, current_rows)
    reconciliation_bytes = csv_bytes(RECONCILIATION_COLUMNS, reconciliation_rows)
    exclusion_bytes = csv_bytes(EXCLUSION_COLUMNS, exclusions)
    build_payload = {
        "builder_sha256": sha256_file(Path(__file__).resolve()),
        "registry_sha256": sha256_bytes(registry_bytes),
        "price_archive_sha256": registry["price_source"]["sha256"],
        "holdings_snapshot_sha256": registry["current_holdings_source"]["sha256"],
        "canonical_price_hashes": [item["canonical_sha256"] for item in instruments],
        "security_master_sha256": sha256_bytes(security_master_bytes),
        "membership_intervals_sha256": sha256_bytes(interval_bytes),
        "current_constituents_sha256": sha256_bytes(current_bytes),
        "snapshot_reconciliation_sha256": sha256_bytes(reconciliation_bytes),
        "snapshot_exclusions_sha256": sha256_bytes(exclusion_bytes),
        "runtime": {
            "python": platform.python_version(),
            "exchange_calendars": xcals.__version__,
            "openpyxl": openpyxl.__version__,
        },
    }
    build_id = sha256_bytes(json.dumps(build_payload, sort_keys=True).encode("utf-8"))
    manifest_path = workspace / output_config["manifest"]
    generated_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    if manifest_path.is_file():
        try:
            old_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if old_manifest.get("build_id") == build_id:
                generated_at = old_manifest["generated_at_utc"]
        except (json.JSONDecodeError, KeyError):
            pass

    manifest = {
        "schema_version": 1,
        "dataset_id": registry["dataset_id"],
        "generated_at_utc": generated_at,
        "builder": "backtest/scripts/build_sp500_history.py",
        "builder_sha256": build_payload["builder_sha256"],
        "registry": str(registry_path.relative_to(workspace)),
        "registry_sha256": sha256_bytes(registry_bytes),
        "build_id": build_id,
        "baseline_status": "validated_for_point_in_time_research_with_documented_gaps",
        "runtime": build_payload["runtime"],
        "price_source": registry["price_source"],
        "current_holdings_source": registry["current_holdings_source"],
        "canonical_price_schema": registry["canonical_price_schema"],
        "summary": summary,
        "snapshot_reconciliation": reconciliation_summary,
        "outputs": {
            "price_directory": output_config["price_directory"],
            "security_master": {"path": output_config["security_master"], "sha256": sha256_bytes(security_master_bytes)},
            "membership_intervals": {"path": output_config["membership_intervals"], "sha256": sha256_bytes(interval_bytes)},
            "current_constituents": {"path": output_config["current_constituents"], "sha256": sha256_bytes(current_bytes)},
            "snapshot_reconciliation": {"path": output_config["snapshot_reconciliation"], "sha256": sha256_bytes(reconciliation_bytes)},
            "snapshot_exclusions": {"path": output_config["snapshot_exclusions"], "sha256": sha256_bytes(exclusion_bytes)},
        },
        "price_files": [
            {
                "instrument_id": item["instrument_id"],
                "path": f"{output_config['price_directory']}/{item['source_member']}",
                "rows": item["price_rows"],
                "start": item["price_start"],
                "end": item["price_end"],
                "source_member_sha256": item["source_member_sha256"],
                "canonical_sha256": item["canonical_sha256"],
                "quality_status": item["quality_status"],
                "member_price_coverage_status": item["member_price_coverage_status"],
            }
            for item in instruments
        ],
        "frozen_existing_outputs": registry["frozen_existing_outputs"],
    }
    quality = {
        "schema_version": 1,
        "generated_at_utc": generated_at,
        "build_id": build_id,
        "summary": summary,
        "reports": reports,
    }
    quality_markdown = render_quality_markdown(manifest, quality, registry).encode("utf-8")

    if failed_instruments:
        shutil.rmtree(price_target_root)
        failed_ids = [item["instrument_id"] for item in failed_instruments[:10]]
        raise ValueError(
            f"{len(failed_instruments)} instruments failed blocking quality checks; "
            f"the previous approved baseline was preserved. Examples: {failed_ids}"
        )

    atomic_write(workspace / output_config["security_master"], security_master_bytes)
    atomic_write(workspace / output_config["membership_intervals"], interval_bytes)
    atomic_write(workspace / output_config["current_constituents"], current_bytes)
    atomic_write(workspace / output_config["snapshot_reconciliation"], reconciliation_bytes)
    atomic_write(workspace / output_config["snapshot_exclusions"], exclusion_bytes)
    atomic_write(
        workspace / output_config["quality_report_json"],
        (json.dumps(quality, ensure_ascii=False, indent=2) + "\n").encode("utf-8"),
    )
    atomic_write(workspace / output_config["quality_report_markdown"], quality_markdown)
    atomic_write(manifest_path, (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    previous_price_root = price_root.parent / f".{price_root.name}.previous"
    if previous_price_root.exists():
        shutil.rmtree(previous_price_root)
    if price_root.exists():
        os.replace(price_root, previous_price_root)
    try:
        os.replace(price_target_root, price_root)
    except BaseException:
        if previous_price_root.exists() and not price_root.exists():
            os.replace(previous_price_root, price_root)
        raise
    if previous_price_root.exists():
        shutil.rmtree(previous_price_root)

    print(
        f"S&P 500 history: instruments={summary['instrument_count']}; "
        f"rows={summary['price_row_count']}; intervals={summary['membership_interval_count']}; "
        f"failed={summary['failed_instruments']}; build_id={build_id}"
    )
    print(
        f"Snapshots: vendor={len(vendor_active)}; official={len(official_active)}; "
        f"official_only={sorted(official_only)}; vendor_only={sorted(vendor_only)}"
    )
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", type=Path, default=WORKSPACE_ROOT)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument(
        "--sample",
        help="Comma-separated instrument IDs; validates them in a temporary directory without changing canonical outputs.",
    )
    args = parser.parse_args()
    workspace = args.workspace.resolve()
    registry_path = args.registry if args.registry.is_absolute() else workspace / args.registry
    sample = {item.strip() for item in args.sample.split(",") if item.strip()} if args.sample else None
    result = build(registry_path.resolve(), workspace, sample)
    if sample is not None:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
